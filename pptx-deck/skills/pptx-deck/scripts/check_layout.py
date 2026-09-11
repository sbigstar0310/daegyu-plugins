#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Measure a deck instead of squinting at it. Exit code 1 means something is wrong.

    python3 scripts/check_layout.py deck.pptx
    python3 scripts/check_layout.py deck.pptx --margins 0.67,9.33,5.30
    python3 scripts/check_layout.py deck.pptx --tokens assets/tokens_white.py
    python3 scripts/check_layout.py deck.pptx --lang-check --strict

Six checks, all of them arithmetic on real coordinates and real font metrics:

  bounds     a shape that leaves the canvas, or crosses the content margins
  overlap    two shapes whose boxes intersect by more than a hairline
  overflow   text that needs more height than its box has
  orphans    a last line holding one or two words, or one wrap away from it
  type       a run below the readable floor, and the type scale actually in use
  gaps       a horizontal band of dead space taller than the budget

Overlap has two exemptions, and you should use them rather than turning the check
off. One shape fully inside another is deliberate: a label on a card, a caption on
a figure. And any shape whose name contains "allow:" is skipped, so a diagram that
really does layer boxes can say so in the file:

    shp.name = "allow: the arrow crosses the band on purpose"

--lang-check is off by default. It looks for the characters that make a deck read
as machine written: the em dash and the middle dot used as a separator.
"""
import argparse
import os
import sys
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import fix_orphans as M  # noqa: E402

EMU = 914400.0
EPS = 0.02          # inches. Two boxes closer than this are touching, not overlapping.
GAP_BUDGET = 1.20   # inches of dead vertical space before a slide looks unfinished.
MIN_PT = 8.0        # nothing below this is readable from the back of a room.
# This is the only file in the skill allowed to contain these two characters,
# because they are what it looks for. Do not "clean" them out.
BANNED = {"—": "em dash", "·": "middle dot separator"}

TABLE = 19  # MSO_SHAPE_TYPE.TABLE


class Report(object):
    def __init__(self):
        self.rows = []

    def add(self, level, check, slide, msg):
        self.rows.append((level, check, slide, msg))

    def errors(self):
        return [r for r in self.rows if r[0] == "ERROR"]

    def warnings(self):
        return [r for r in self.rows if r[0] == "WARN"]


def box(shape):
    """(left, top, right, bottom) in inches."""
    l = (shape.left or 0) / EMU
    t = (shape.top or 0) / EMU
    return l, t, l + (shape.width or 0) / EMU, t + (shape.height or 0) / EMU


def is_plain_textbox(shape):
    """A text box with no fill and no outline is mostly air: only the glyphs are
    on the slide, so only the glyphs can collide with anything."""
    if not shape.has_text_frame or shape.shape_type == TABLE:
        return False
    try:
        if shape.fill.type not in (None, 5):      # 5 = MSO_FILL.BACKGROUND
            return False
        if shape.line.fill.type not in (None, 5):
            return False
    except (NotImplementedError, ValueError):
        return False
    return True


def ink_box(shape, family):
    """The box the shape actually covers. For a plain text box that is the text
    extent, placed inside the frame according to its alignment and anchor. Falls
    back to the geometric box when the font cannot be measured."""
    l, t, r, b = box(shape)
    if not is_plain_textbox(shape):
        return l, t, r, b
    tf = shape.text_frame
    if not tf.text.strip():
        return l, t, l, t
    pad_l = (tf.margin_left or 0) / EMU
    pad_r = (tf.margin_right or 0) / EMU
    pad_t = (tf.margin_top or 0) / EMU
    pad_b = (tf.margin_bottom or 0) / EMU
    inner_w = (r - l) - pad_l - pad_r
    widest = 0.0
    height = 0.0
    centred = False
    right_aligned = False
    for p in tf.paragraphs:
        if not p.runs or p.runs[0].font.size is None:
            return l, t, r, b
        size = p.runs[0].font.size.pt
        toks = M.tokens(p, family)
        if not toks:
            continue
        ls = M.lines(toks, size, M.para_width(shape, p))
        if ls is None:
            return l, t, r, b
        widest = max(widest, max(ls) / 72.0)
        spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
        after = p.space_after.pt if p.space_after is not None else 0
        height += (len(ls) * size * spacing + after) / 72.0
        if p.alignment is not None:
            centred = centred or str(p.alignment).startswith("CENTER")
            right_aligned = right_aligned or str(p.alignment).startswith("RIGHT")
    if widest <= 0:
        return l, t, r, b
    widest = min(widest, inner_w)
    height = min(height, (b - t) - pad_t - pad_b) if (b - t) > pad_t + pad_b else height
    if centred:
        x0 = l + pad_l + (inner_w - widest) / 2.0
    elif right_aligned:
        x0 = r - pad_r - widest
    else:
        x0 = l + pad_l
    anchor = str(tf.vertical_anchor or "")
    inner_h = (b - t) - pad_t - pad_b
    if "MIDDLE" in anchor:
        y0 = t + pad_t + max(0.0, (inner_h - height) / 2.0)
    elif "BOTTOM" in anchor:
        y0 = b - pad_b - height
    else:
        y0 = t + pad_t
    return x0, y0, x0 + widest, y0 + height


def visible_shapes(slide):
    for shape in slide.shapes:
        if shape.width is None or shape.height is None:
            continue
        if shape.width <= 0 or shape.height <= 0:
            continue
        # An empty text box occupies no ink and is not worth a complaint.
        if shape.has_text_frame and not shape.text_frame.text.strip() \
                and shape.shape_type != TABLE and not shape.fill.type:
            continue
        yield shape


def infer_margins(prs):
    """The margins the deck actually uses: the most common left edge, the most
    common right edge, and the lowest bottom that is not the page number."""
    lefts, rights, bottoms = Counter(), Counter(), []
    for slide in prs.slides:
        for shape in visible_shapes(slide):
            l, t, r, b = box(shape)
            lefts[round(l, 2)] += 1
            rights[round(r, 2)] += 1
            if (shape.height or 0) / EMU > 0.25:
                bottoms.append(b)
    ml = lefts.most_common(1)[0][0] if lefts else 0.5
    mr = rights.most_common(1)[0][0] if rights else prs.slide_width / EMU - 0.5
    bottoms.sort()
    cb = bottoms[int(len(bottoms) * 0.98)] if bottoms else prs.slide_height / EMU - 0.3
    return ml, mr, round(cb, 2)


def load_tokens(path):
    import importlib.util
    spec = importlib.util.spec_from_file_location("deck_tokens", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ---------------------------------------------------------------- the checks
def check_bounds(prs, rep, ml, mr, cb):
    W = prs.slide_width / EMU
    H = prs.slide_height / EMU
    for i, slide in enumerate(prs.slides):
        n = i + 1
        for shape in visible_shapes(slide):
            l, t, r, b = box(shape)
            name = shape.name
            if l < -0.01 or t < -0.01 or r > W + 0.01 or b > H + 0.01:
                rep.add("ERROR", "bounds", n,
                        "%s leaves the canvas: %.2f %.2f to %.2f %.2f, canvas is %.2f x %.2f"
                        % (name, l, t, r, b, W, H))
            elif l < ml - 0.02 or r > mr + 0.02:
                rep.add("WARN", "bounds", n,
                        "%s crosses the side margin: %.2f to %.2f, margins are %.2f to %.2f"
                        % (name, l, r, ml, mr))
            elif b > cb + 0.02 and t < cb:
                rep.add("WARN", "bounds", n,
                        "%s ends at %.2f, below the content line %.2f" % (name, b, cb))


def check_overlap(prs, rep, eps, family):
    for i, slide in enumerate(prs.slides):
        n = i + 1
        shapes = [s for s in visible_shapes(slide) if "allow:" not in (s.name or "")]
        boxes = [ink_box(s, family) for s in shapes]
        for a in range(len(shapes)):
            for c in range(a + 1, len(shapes)):
                la, ta, ra, ba = boxes[a]
                lb, tb, rb, bb = boxes[c]
                ox = min(ra, rb) - max(la, lb)
                oy = min(ba, bb) - max(ta, tb)
                if ox <= eps or oy <= eps:
                    continue
                # Containment is deliberate: a label on a card, a caption on a figure.
                inside = ((la >= lb - eps and ra <= rb + eps and ta >= tb - eps and ba <= bb + eps)
                          or (lb >= la - eps and rb <= ra + eps and tb >= ta - eps and bb <= ba + eps))
                if inside:
                    continue
                rep.add("ERROR", "overlap", n,
                        "%s and %s overlap by %.2f x %.2f in"
                        % (shapes[a].name, shapes[c].name, ox, oy))


def check_overflow(prs, rep, family):
    for i, slide in enumerate(prs.slides):
        n = i + 1
        for shape in visible_shapes(slide):
            if not shape.has_text_frame or shape.shape_type == TABLE:
                continue
            tf = shape.text_frame
            if tf.word_wrap is False:
                continue
            need = ((tf.margin_top or 0) + (tf.margin_bottom or 0)) / EMU
            measured = False
            for p in tf.paragraphs:
                if not p.runs or p.runs[0].font.size is None:
                    continue
                size = p.runs[0].font.size.pt
                toks = M.tokens(p, family)
                if not toks:
                    continue
                ls = M.lines(toks, size, M.para_width(shape, p))
                if ls is None:
                    continue
                measured = True
                spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
                after = p.space_after.pt if p.space_after is not None else 0
                need += (len(ls) * size * spacing + after) / 72.0
            have = (shape.height or 0) / EMU
            if measured and need > have + 0.06:
                rep.add("ERROR", "overflow", n,
                        "%s needs %.2f in of height and has %.2f" % (shape.name, need, have))


def check_orphans(prs, rep, family, min_ratio, fill, include_first):
    for n, text, kind in M.report(prs, family, min_ratio, not include_first, fill):
        rep.add("WARN" if kind == "at risk" else "ERROR", "orphans", n,
                "%s: %s" % (kind, text[:64]))


def check_type(prs, rep, min_pt):
    sizes = Counter()
    small = Counter()
    example = {}
    for i, slide in enumerate(prs.slides):
        n = i + 1
        for shape in visible_shapes(slide):
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size is None or not r.text.strip():
                        continue
                    pt = round(r.font.size.pt, 2)
                    sizes[pt] += len(r.text)
                    if pt < min_pt:
                        small[(n, pt)] += 1
                        example.setdefault((n, pt), r.text.strip()[:36])
    for (n, pt), count in sorted(small.items()):
        rep.add("WARN", "type", n,
                "%d run%s at %.1f pt, below the floor of %.1f, such as %r"
                % (count, "" if count == 1 else "s", pt, min_pt, example[(n, pt)]))
    return sizes


def check_gaps(prs, rep, ml, mr, cb, budget):
    top_of_body = 1.10
    for i, slide in enumerate(prs.slides):
        n = i + 1
        spans = []
        for shape in visible_shapes(slide):
            l, t, r, b = box(shape)
            if r < ml - 0.5 or l > mr + 0.5:
                continue
            if b <= top_of_body or t >= cb:
                continue
            spans.append((max(t, top_of_body), min(b, cb)))
        if not spans:
            continue
        spans.sort()
        merged = [list(spans[0])]
        for t, b in spans[1:]:
            if t <= merged[-1][1] + 0.001:
                merged[-1][1] = max(merged[-1][1], b)
            else:
                merged.append([t, b])
        edges = [top_of_body] + [x for m in merged for x in m] + [cb]
        worst = 0.0
        at = None
        for k in range(0, len(edges) - 1, 2):
            gap = edges[k + 1] - edges[k]
            if gap > worst:
                worst, at = gap, edges[k]
        if worst > budget:
            rep.add("WARN", "gaps", n,
                    "%.2f in of dead space starting at y %.2f, budget is %.2f"
                    % (worst, at, budget))


def check_language(prs, rep):
    for i, slide in enumerate(prs.slides):
        n = i + 1
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
        for t in texts:
            for ch, what in BANNED.items():
                if ch in t:
                    where = t[max(0, t.index(ch) - 28):t.index(ch) + 28].replace("\n", " ")
                    rep.add("WARN", "language", n, "%s in: %s" % (what, where))


# ---------------------------------------------------------------- driver
def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("deck")
    ap.add_argument("--tokens", help="a tokens.py to read ML, MR and CONTENT_BOTTOM from")
    ap.add_argument("--margins", help="L,R,BOTTOM in inches, overriding everything else")
    ap.add_argument("--font", default="Inter", help="family to measure with")
    ap.add_argument("--eps", type=float, default=EPS, help="overlap tolerance in inches")
    ap.add_argument("--gap-budget", type=float, default=GAP_BUDGET)
    ap.add_argument("--min-pt", type=float, default=MIN_PT)
    ap.add_argument("--min-ratio", type=float, default=0.34)
    ap.add_argument("--fill-risk", type=float, default=M.AT_RISK_FILL)
    ap.add_argument("--include-first", action="store_true")
    ap.add_argument("--lang-check", action="store_true",
                    help="also flag em dashes and middle dots. Off by default.")
    ap.add_argument("--only", help="run only these checks, comma separated")
    ap.add_argument("--strict", action="store_true", help="warnings fail too")
    a = ap.parse_args()

    from pptx import Presentation
    prs = Presentation(a.deck)

    if a.margins:
        ml, mr, cb = [float(x) for x in a.margins.split(",")]
        source = "given"
    elif a.tokens:
        T = load_tokens(a.tokens)
        ml, mr, cb = T.ML, T.MR, T.CONTENT_BOTTOM
        source = os.path.basename(a.tokens)
    else:
        ml, mr, cb = infer_margins(prs)
        source = "inferred from the deck"

    only = {x.strip() for x in a.only.split(",")} if a.only else None

    def run(name):
        return only is None or name in only

    rep = Report()
    if run("bounds"):
        check_bounds(prs, rep, ml, mr, cb)
    if run("overlap"):
        check_overlap(prs, rep, a.eps, a.font)
    if run("overflow"):
        check_overflow(prs, rep, a.font)
    if run("orphans"):
        check_orphans(prs, rep, a.font, a.min_ratio, a.fill_risk, a.include_first)
    sizes = check_type(prs, rep, a.min_pt) if run("type") else Counter()
    if run("gaps"):
        check_gaps(prs, rep, ml, mr, cb, a.gap_budget)
    if a.lang_check and run("language"):
        check_language(prs, rep)

    print("deck     %s, %d slides, canvas %.2f x %.2f in"
          % (os.path.basename(a.deck), len(prs.slides),
             prs.slide_width / EMU, prs.slide_height / EMU))
    print("margins  %.2f to %.2f, content bottom %.2f  (%s)" % (ml, mr, cb, source))
    if sizes:
        scale = ", ".join("%.1f" % s for s, _n in sorted(sizes.items(), reverse=True))
        print("type     %d sizes in use: %s" % (len(sizes), scale))
        if len(sizes) > 8:
            print("         that is a lot. A scale the audience can read has five or six.")
    print("")

    by_check = Counter(r[1] for r in rep.rows)
    for level in ("ERROR", "WARN"):
        rows = [r for r in rep.rows if r[0] == level]
        for _l, check, slide, msg in sorted(rows, key=lambda r: (r[1], r[2])):
            print("%-5s %-9s slide %-3d %s" % (level, check, slide, msg))
    if rep.rows:
        print("")
    print("%d error, %d warning%s" % (len(rep.errors()), len(rep.warnings()),
                                      "  (" + ", ".join("%s %d" % (k, v)
                                                        for k, v in sorted(by_check.items())) + ")"
                                      if by_check else ""))
    if rep.errors() or (a.strict and rep.warnings()):
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
