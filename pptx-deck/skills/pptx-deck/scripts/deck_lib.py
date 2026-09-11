#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""The python-pptx mechanics, and nothing else.

Every function here works around something python-pptx cannot do, or does wrong
by default. There are no coordinates, no palette, no slide archetypes and no house
style: those belong to your deck's tokens, because they are the part that changes
with the reference. Import this, import your tokens, write build().

    import sys, os
    sys.path.insert(0, "path/to/skills/pptx-deck/scripts")
    import deck_lib as D
    D.FONT = "Inter"

What each one is for:

    blank        a slide with no placeholders, which is the only sane starting point
    textbox      a text box with zero insets and per-run weight and colour
    text         the one line version
    shape        an autoshape with the theme fill and the drop shadow removed
    picture      returns the size it placed, so a caption can sit under it
    connector    a straight line, with a real arrow head
    table        no header fill, hairlines only where you ask for them
    cell_border  per side borders, which python-pptx has no API for
    highlight    a marker behind a run, which python-pptx has no API for
    bullets      real bullet paragraphs with a hanging indent
    notes        speaker notes
    hide         keep a slide in the file and out of the slideshow and the PDF
    fix_orphans  no paragraph ends with one word alone
"""
import os

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONT = "Inter"      # set this once, after importing
BLACK = RGBColor(0, 0, 0)

EMU = 914400


# ---------------------------------------------------------------- slides
def canvas(prs, w_in, h_in):
    prs.slide_width = Inches(w_in)
    prs.slide_height = Inches(h_in)
    return prs


def blank(prs):
    """Layout 6 is the blank one. Any other layout drops placeholders on the slide
    that you then have to fight."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def hide(slide, hidden=True):
    """Kept in the file, excluded from the slideshow and from the exported PDF,
    and still reachable during the talk by typing its number. This is how an
    appendix should be handled: hide it, do not cut it."""
    slide._element.set("show", "0" if hidden else "1")
    return slide


def is_hidden(slide):
    return slide._element.get("show") == "0"


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt.strip()
    return slide


# ---------------------------------------------------------------- text
def highlight(run, color):
    """A marker behind one run. There is no API for this, and the element must come
    BEFORE a:latin in the run properties or PowerPoint drops it without a word.
    color is a hex string with no hash, for example "E6E6E6"."""
    rPr = run._r.get_or_add_rPr()
    el = etree.Element(qn("a:highlight"))
    etree.SubElement(el, qn("a:srgbClr")).set("val", color.lstrip("#"))
    latin = rPr.find(qn("a:latin"))
    if latin is not None:
        latin.addprevious(el)
    else:
        rPr.append(el)
    return run


def textbox(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT,
            font=None, inset=0.0):
    """paras is a list of dicts. Each one takes either

        t     the whole paragraph as one run
        runs  [(text, bold, color_or_None, marker_hex_or_None)] for mixed weight

    and optionally size, bold, color, after, line, align.

    Insets are zeroed by default, because a default text box carries 0.1 in of
    padding on the left and right that silently eats into your measured width.
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(inset)
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", align)
        p.space_before = Pt(spec.get("before", 0))
        p.space_after = Pt(spec.get("after", 4))
        p.line_spacing = spec.get("line", 1.12)
        runs = spec.get("runs")
        if runs is None:
            runs = [(spec.get("t", ""), spec.get("bold", False), spec.get("color"), None)]
        for item in runs:
            item = tuple(item) + (None,) * (4 - len(item))
            body, bold, color, mark = item[:4]
            if not body:
                continue
            r = p.add_run()
            r.text = body
            r.font.size = Pt(spec.get("size", 12))
            r.font.bold = bool(bold)
            r.font.color.rgb = color or spec.get("color") or BLACK
            r.font.name = font or FONT
            if mark:
                highlight(r, mark if isinstance(mark, str) else "E6E6E6")
    return box


def text(slide, x, y, w, h, body, size=12, bold=False, color=None,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=None):
    return textbox(slide, x, y, w, h,
                   [dict(t=body, size=size, bold=bold, color=color, after=0)],
                   anchor=anchor, align=align, font=font)


def bullets(slide, x, y, w, h, entries, size=12, after=6, indent=0.2, font=None,
            color=None):
    """entries is [(lead, body)] or [body]. A real hanging indent, so the second
    line of a wrapped bullet lines up under the first word and not under the dot."""
    paras = []
    for e in entries:
        if isinstance(e, (tuple, list)) and len(e) == 2:
            lead, body = e
            paras.append(dict(runs=[(lead + ": ", True, color, None),
                                    (body, False, color, None)], size=size, after=after))
        else:
            paras.append(dict(t=e if isinstance(e, str) else e[0], size=size,
                              after=after, color=color))
    box = textbox(slide, x, y, w, h, paras, font=font)
    for p in box.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(indent))))
        pPr.set("indent", str(-int(Inches(indent))))
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
    return box


def shape_text(shp, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER,
               margin=0.08, font=None):
    """Text inside an autoshape. Note that a run centred vertically inside a shape
    does not land in the same place in LibreOffice and in PowerPoint, so never put
    a lone digit in an oval and trust the render: draw numbered markers as images."""
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spec.get("after", 2))
        r = p.add_run()
        r.text = spec.get("t", "")
        r.font.size = Pt(spec.get("size", 11))
        r.font.bold = spec.get("bold", False)
        r.font.color.rgb = spec.get("color") or BLACK
        r.font.name = font or FONT
    return shp


# ---------------------------------------------------------------- geometry
def shape(slide, x, y, w, h, fill=None, line=None, lw=0.75,
          kind=MSO_SHAPE.RECTANGLE, radius=None, name=None):
    """An autoshape with the theme's fill and its drop shadow removed. Without
    this every rectangle arrives tinted with the theme accent and wearing a
    shadow, and the deck grows a colour it does not have."""
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)
    if name:
        shp.name = name
    return shp


def picture(slide, path, x, y, w=None, h=None, name=None):
    """Returns (width, height) in inches, so the caption below it can be anchored
    to where the picture actually ends instead of where you guessed it would."""
    from PIL import Image
    iw, ih = Image.open(path).size
    if w is None and h is None:
        raise ValueError("give picture() a width or a height")
    if w is None:
        w = h * iw / ih
    if h is None:
        h = w * ih / iw
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        p.name = name
    return w, h


def connector(slide, x1, y1, x2, y2, color=None, lw=1.0, arrow=True, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color or BLACK
    c.line.width = Pt(lw)
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    if name:
        c.name = name
    return c


# ---------------------------------------------------------------- tables
NO_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def cell_border(cell, sides="B", color="CFCFCF", w_pt=0.75):
    """Per side borders. python-pptx has no API for these, and getting them to
    draw took three tries. Three things have to be right.

    All four elements must be written, every time. A side you leave out is not
    borderless: the table style's own border draws there instead. The unwanted
    sides get a:noFill, which is what actually turns them off.

    They must be in the order L, R, T, B, and they must be the FIRST children of
    a:tcPr, ahead of the fill. Appending after the fill is invalid and both
    PowerPoint and LibreOffice drop the borders without saying anything.

    And the table itself must carry the no-style id, or the style repaints over
    all of this. table() below does that part.

    sides is any of "LRTB", or "" for a cell with no border at all.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    for i, side in enumerate(["L", "R", "T", "B"]):
        tag = qn("a:ln" + side)
        old = tcPr.find(tag)
        if old is not None:
            tcPr.remove(old)
        ln = etree.Element(tag)
        ln.set("w", str(int(w_pt * 12700)))
        if side in sides.upper():
            fill = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr")).set("val", color.lstrip("#"))
        else:
            etree.SubElement(ln, qn("a:noFill"))
        tcPr.insert(i, ln)
    return cell


def table(slide, x, y, header, rows, col_w, row_h=0.30, hdr_h=0.28, size=11,
          font=None, color=None, rule="CFCFCF", rule_pt=0.5, sides="B",
          fill=None, first_col_bold=False, name=None):
    """A table with no header fill: hairlines only where you ask for them, which
    is what an academic reference does. The default is one line under every row.

    Returns the y of the table's requested bottom. Treat that as a request, not a
    fact: if a cell wraps one line further than you expected, PowerPoint grows
    that row and everything below the table moves down. Measure the real bottom on
    the render before you place anything under it.
    """
    from pptx.dml.color import RGBColor as _RGB
    white = _RGB(0xFF, 0xFF, 0xFF)
    n_rows = len(rows) + 1
    g = slide.shapes.add_table(n_rows, len(col_w), Inches(x), Inches(y),
                               Inches(sum(col_w)), Inches(hdr_h + row_h * len(rows)))
    tbl = g.table
    tbl.first_row = False       # otherwise the theme paints a banner behind row 1
    tbl.horz_banding = False
    # Neutralise the built in table style, which would otherwise repaint the
    # fills and the borders set below.
    tblPr = tbl._tbl.tblPr
    for el in list(tblPr):
        if el.tag == qn("a:tableStyleId"):
            tblPr.remove(el)
    etree.SubElement(tblPr, qn("a:tableStyleId")).text = NO_TABLE_STYLE

    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    data = [list(header)] + [list(r) for r in rows]
    for r in range(n_rows):
        tbl.rows[r].height = Inches(hdr_h if r == 0 else row_h)
        for c in range(len(col_w)):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0 if c == 0 else 0.10)
            cell.margin_right = Inches(0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill or white
            cell_border(cell, sides, rule, rule_pt)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = 1.12
            run = p.add_run()
            run.text = str(data[r][c])
            run.font.name = font or FONT
            run.font.size = Pt(size)
            run.font.color.rgb = color or BLACK
            run.font.bold = (r == 0) or (c == 0 and first_col_bold)
    if name:
        g.name = name
    return y + hdr_h + row_h * len(rows)


# ---------------------------------------------------------------- the last pass
def fix_orphans(prs, family=None, **kw):
    """Run this once, on the finished deck, just before saving. It shrinks a font
    or narrows a box by a few percent so that no paragraph ends with one word
    alone. Returns how many it changed."""
    import sys
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import fix_orphans as F
    return F.fix(prs, family or FONT, **kw)
