#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Render a deck the way a projector will: LibreOffice to PDF, then PDF to PNG.

    python3 scripts/render_real.py deck.pptx                 # -> render/s01.png ...
    python3 scripts/render_real.py deck.pptx --dpi 200 -o build/render
    python3 scripts/render_real.py deck.pptx --slides 8,13   # re-render two pages
    python3 scripts/render_real.py deck.pptx --contact-sheet # and render/contact.png
    python3 scripts/render_real.py deck.pptx --pdf-out docs/deck.pdf   # the deliverable
    python3 scripts/render_real.py deck.pptx --skip-hidden --pdf-out handout.pdf

This is the only render you are allowed to judge line breaks from. The matplotlib
previewer uses a proxy font and approximates wrapping, so it can tell you where a
box sits and never where a line ends.

Hidden slides are rendered too. A render made for checking shows everything the
file holds, so the appendix gets its layout checked like every other slide, and
page N is slide N. The hidden ones are listed. --skip-hidden renders what the
slideshow shows instead, which is what a handout PDF wants.

--pdf-out keeps the PDF this run rendered from, hidden slides included unless
--skip-hidden, so the file you deliver is the file you checked. --contact-sheet
tiles every slide it wrote into contact.png, at SHEET_DPI whatever --dpi says,
each labelled with its slide number.

Three traps it handles for you.

Skipped hidden slides shift the pages, so with --skip-hidden page N is not slide
N. The map is printed, and the PNG for a slide is named by its slide number, not
by its page number.

A render older than the deck is not evidence. Any PNG or PDF left from a previous
run is deleted before this one starts, and the deck's own mtime is printed.

LibreOffice substitutes a missing font in silence. Every family the deck names is
resolved against the font directories first, and a missing one is a hard error,
because every line break in the output would be fiction. A missing Bold whose
Regular is there is only a warning: LibreOffice strokes the Regular glyphs, which
keeps their widths, so the render is still true to this machine.

LibreOffice alone widens every gap between Hangul and Latin text. When a python
that can import uno is found, the export goes through uno_pdf.py, which turns that
off; otherwise it is the plain command line export. The route that ran is printed.
"""
import argparse
import glob
import os
import shutil
import subprocess
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from fix_orphans import font_file  # noqa: E402
from uno_pdf import find_python as find_uno_python  # noqa: E402

UNO_PDF = os.path.join(os.path.dirname(os.path.abspath(__file__)), "uno_pdf.py")

# LibreOffice leaves hidden slides out of a PDF unless told otherwise. JSON filter
# options on the command line need LibreOffice 7.4 or newer.
PDF_WITH_HIDDEN = 'pdf:impress_pdf_Export:{"ExportHiddenSlides":{"type":"boolean","value":"true"}}'

SHEET_DPI = 120    # the whole deck in one look, as SKILL.md reviews it
SHEET_COLUMNS = 4
SHEET_PAD = 16     # pixels between tiles, and under each for its label

SOFFICE = [
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice", "/usr/local/bin/soffice", "/snap/bin/libreoffice",
]


def find_soffice():
    for p in SOFFICE:
        if os.path.exists(p):
            return p
    return shutil.which("soffice") or shutil.which("libreoffice")


def families(prs):
    """Every (family, bold) named by a run in the deck, most used first."""
    from collections import Counter
    c = Counter()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.name:
                        c[(r.font.name, bool(r.font.bold))] += len(r.text)
    return [f for f, _n in c.most_common()]


def is_hidden(slide):
    # An OOXML boolean may be written "false" as well as "0", and LibreOffice
    # hides the slide either way.
    return slide._element.get("show") in ("0", "false")


def slide_map(prs, skip_hidden):
    """[(slide_number, page_number or None)]. With hidden slides exported, page N
    is slide N. With them skipped, a hidden slide has no page and the pages after
    it shift up."""
    out = []
    page = 0
    for i, slide in enumerate(prs.slides):
        if skip_hidden and is_hidden(slide):
            out.append((i + 1, None))
        else:
            page += 1
            out.append((i + 1, page))
    return out


def contact_sheet(doc, pages, hidden):
    """pages is [(slide_number, page_number)]. One image, SHEET_COLUMNS across,
    each slide on a grey ground so a white slide keeps its edge."""
    from PIL import Image, ImageDraw, ImageFont
    tiles = []
    for slide_no, page_no in pages:
        pix = doc[page_no - 1].get_pixmap(dpi=SHEET_DPI)
        tiles.append((slide_no, Image.frombytes("RGB", (pix.width, pix.height), pix.samples)))
    w = max(im.width for _n, im in tiles)
    h = max(im.height for _n, im in tiles)
    label = 2 * SHEET_PAD
    cols = min(SHEET_COLUMNS, len(tiles))
    rows = (len(tiles) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * (w + SHEET_PAD) + SHEET_PAD,
                              rows * (h + label + SHEET_PAD) + SHEET_PAD), "#DDDDDD")
    draw = ImageDraw.Draw(sheet)
    try:
        font = ImageFont.load_default(size=label - 8)
    except TypeError:  # Pillow before 10.1 has only a small bitmap font
        font = ImageFont.load_default()
    for i, (slide_no, im) in enumerate(tiles):
        r, c = divmod(i, cols)
        x = SHEET_PAD + c * (w + SHEET_PAD)
        y = SHEET_PAD + r * (h + label + SHEET_PAD)
        sheet.paste(im, (x, y))
        draw.text((x, y + h + 4), "%d%s" % (slide_no, "  hidden" if slide_no in hidden else ""),
                  fill="black", font=font)
    return sheet


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("deck")
    ap.add_argument("-o", "--out", default="render", help="directory for the PNGs")
    ap.add_argument("--dpi", type=int, default=120, help="120 to review layout, 200 to read text")
    ap.add_argument("--slides", help="only write these slide numbers, comma separated")
    ap.add_argument("--keep-pdf", action="store_true")
    ap.add_argument("--pdf-out", help="also write the checked PDF to this path")
    ap.add_argument("--contact-sheet", action="store_true",
                    help="also tile every slide written into contact.png at %d dpi" % SHEET_DPI)
    ap.add_argument("--skip-hidden", action="store_true",
                    help="leave hidden slides out, as the slideshow does. For a handout PDF.")
    ap.add_argument("--allow-missing-font", action="store_true",
                    help="render anyway. Line breaks in the output will be wrong.")
    a = ap.parse_args()

    from pptx import Presentation
    prs = Presentation(a.deck)

    soffice = find_soffice()
    if not soffice:
        sys.stderr.write(
            "LibreOffice is not installed, and there is no other way to see the real\n"
            "line breaks. Install it:\n"
            "    brew install --cask libreoffice        # macOS\n"
            "    sudo apt-get install -y libreoffice    # Debian or Ubuntu\n")
        return 2

    used = families(prs)
    missing = []
    for f, _bold in used:
        if f not in missing and font_file(f, False) is None:
            missing.append(f)
    # A bold with no Bold file but a Regular is drawn by stroking the Regular, which
    # keeps its widths: the render is real, a machine with the Bold may differ.
    for f, bold in used:
        if bold and f not in missing and font_file(f, True) is None:
            sys.stderr.write(
                'WARNING: "%s" has no Bold file, so LibreOffice will synthesise the bold\n'
                "from the Regular. Its line breaks are real here; on a machine with the\n"
                "real Bold a proportional family can set up to ~7%% wider.\n" % f)
    if missing:
        msg = ("the deck asks for %s and no file for it was found.\n"
               "LibreOffice will substitute another font without saying so, and every\n"
               "line break below would be fiction. Install the font, or copy it into\n"
               "/Applications/LibreOffice.app/Contents/Resources/fonts/truetype\n"
               % ", ".join(missing))
        if not a.allow_missing_font:
            sys.stderr.write(msg)
            return 2
        sys.stderr.write("WARNING: " + msg)

    out = os.path.abspath(a.out)
    os.makedirs(out, exist_ok=True)
    # Not `rm -f DIR/*.png`: zsh aborts the whole command on an unmatched glob.
    for old in glob.glob(os.path.join(out, "s*.png")) + glob.glob(os.path.join(out, "contact.png")):
        os.remove(old)

    deck = os.path.abspath(a.deck)
    print("deck    %s" % deck)
    print("        modified %s" % time.strftime("%Y-%m-%d %H:%M:%S",
                                                time.localtime(os.path.getmtime(deck))))

    pdf = os.path.join(out, os.path.splitext(os.path.basename(deck))[0] + ".pdf")
    # LibreOffice can fail to write and still exit 0, and an old PDF in its place
    # would be rendered as if it were this deck.
    if os.path.exists(pdf):
        os.remove(pdf)
    route = "command line, autospace on"
    uno_python = find_uno_python(soffice)
    if uno_python:
        cmd = [uno_python, UNO_PDF, soffice, deck, pdf] + ([] if a.skip_hidden else ["--hidden"])
        try:
            err = subprocess.run(cmd, capture_output=True, text=True, timeout=300).stderr
        except subprocess.TimeoutExpired:
            err = "timed out after 300 s"
        if os.path.exists(pdf):
            route = "uno, autospace off"
        else:
            sys.stderr.write("WARNING: the uno export wrote no PDF (%s), so the command line "
                             "export runs.\n" % (err.strip().splitlines() or ["no message"])[-1])
    if not os.path.exists(pdf):
        target = "pdf" if a.skip_hidden else PDF_WITH_HIDDEN
        r = subprocess.run([soffice, "--headless", "--convert-to", target, "--outdir", out, deck],
                           capture_output=True, text=True)
    if not os.path.exists(pdf):
        sys.stderr.write("LibreOffice produced no PDF.\n%s\n%s\n" % (r.stdout, r.stderr))
        sys.stderr.write("If PowerPoint has the deck open, a ~$ lock file is in the way.\n")
        return 2

    import pymupdf
    doc = pymupdf.open(pdf)

    hidden = [i + 1 for i, slide in enumerate(prs.slides) if is_hidden(slide)]
    skipped = a.skip_hidden
    if not skipped and hidden and doc.page_count == len(prs.slides) - len(hidden):
        # An older LibreOffice ignores the filter option. Naming page N slide N
        # would put every slide after the first hidden one under the wrong name.
        sys.stderr.write("WARNING: LibreOffice left the hidden slides out anyway, so they are\n"
                         "not rendered. Exporting them needs LibreOffice 7.4 or newer.\n")
        skipped = True
    smap = slide_map(prs, skipped)
    rendered = [s for s, p in smap if p is not None]
    if doc.page_count != len(rendered):
        sys.stderr.write("WARNING: %d pages but %d slides to render. The map below is a guess.\n"
                         % (doc.page_count, len(rendered)))

    want = None
    if a.slides:
        want = {int(x) for x in a.slides.replace(" ", "").split(",") if x}

    written = []
    for slide_no, page_no in smap:
        if page_no is None or page_no > doc.page_count or (want and slide_no not in want):
            continue
        path = os.path.join(out, "s%02d.png" % slide_no)
        doc[page_no - 1].get_pixmap(dpi=a.dpi).save(path)
        written.append((slide_no, page_no))
    if a.contact_sheet and written:
        contact_sheet(doc, written, hidden).save(os.path.join(out, "contact.png"))
    page_count = doc.page_count
    doc.close()
    dest = os.path.abspath(a.pdf_out) if a.pdf_out else None
    if dest and dest != pdf:
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        shutil.copyfile(pdf, dest)
    if not a.keep_pdf and dest != pdf:
        os.remove(pdf)

    print("export  %s" % route)
    print("slides  %d, visible %d, hidden %d" % (len(smap), len(smap) - len(hidden), len(hidden)))
    print("pages   %d at %d dpi, wrote %d PNG into %s" % (page_count, a.dpi, len(written), out))
    if a.contact_sheet and written:
        print("sheet   contact.png, %d slides at %d dpi" % (len(written), SHEET_DPI))
    if dest:
        print("pdf     %s" % dest)
    if not hidden:
        print("no hidden slides, so page N is slide N")
    elif not skipped:
        print("hidden  %s, rendered like the rest, so page N is slide N."
              % ", ".join(str(s) for s in hidden))
        print("        The slideshow skips them. --skip-hidden renders it that way.")
    else:
        print("\nslide to page map. Hidden slides are in the file but not in the PDF,")
        print("so do not look for slide %d on page %d." % (hidden[0], hidden[0]))
        rows = ["%d>%s" % (s, p if p else "hidden") for s, p in smap]
        for i in range(0, len(rows), 8):
            print("  " + "  ".join(rows[i:i + 8]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
