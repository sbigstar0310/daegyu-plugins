#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""What is missing, and the one command that installs it.

    python3 scripts/preflight.py
    python3 scripts/preflight.py --font Inter --deck deck.pptx

Run this before you promise the user anything. Each row is either present, or
missing with the command that fixes it. Nothing is installed for you: print the
list, tell the user what the deck loses without each one, and let them choose.

Required means the deck cannot be built or cannot be checked. Optional means one
capability is unavailable and the deck is still buildable without it.
"""
import argparse
import importlib.util
import os
import platform
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

BREW = shutil.which("brew") is not None
APT = shutil.which("apt-get") is not None


def pip_cmd(pkg):
    return "%s -m pip install %s" % (os.path.basename(sys.executable), pkg)


def app_cmd(brew, apt):
    if BREW:
        return brew
    if APT:
        return apt
    return brew + "   (or your platform's package manager)"


PY_MODULES = [
    ("pptx", "python-pptx", True, "build or read a .pptx at all"),
    ("PIL", "pillow", True, "measure line breaks and place images"),
    ("lxml", "lxml", True, "write the XML python-pptx has no API for"),
    ("pymupdf", "pymupdf", True, "turn the rendered PDF into PNG you can look at"),
]

# The image plugin takes any one of these, each from its provider's key page.
# OpenRouter is the gateway to all of them.
IMAGE_KEYS = {
    "OPENROUTER_API_KEY": "https://openrouter.ai/keys  (credits: https://openrouter.ai/settings/credits)",
    "GEMINI_API_KEY": "https://aistudio.google.com/apikey",
    "OPENAI_API_KEY": "https://platform.openai.com/api-keys",
    "XAI_API_KEY": "https://console.x.ai",
    "GROK_API_KEY": "https://console.x.ai",
}
# Where the key lives by default. A project's .env holds the project's own keys,
# and an image key there bills the wrong account.
USER_KEY_FILE = "~/.config/pptx-deck/image.env"
# Describes the key it is called with, and costs nothing.
OPENROUTER_KEY_URL = "https://openrouter.ai/api/v1/key"

SKILL_PLUGINS = [
    ("claude-image-generation", "hex-plugins", "hex/claude-marketplace",
     "generate the diagrams and illustrations a slide needs when the paper has none"),
    ("codex", "openai-codex", "openai/codex-plugin-cc",
     "hand a slide to a second model for a redesign or an adversarial check"),
]


def find_soffice():
    for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
              "/usr/bin/soffice", "/usr/local/bin/soffice", "/snap/bin/libreoffice"):
        if os.path.exists(p):
            return p
    return shutil.which("soffice") or shutil.which("libreoffice")


def plugin_installed(name):
    """True if a plugin of this name is in the Claude Code cache."""
    root = os.path.expanduser("~/.claude/plugins/cache")
    if not os.path.isdir(root):
        return False
    for market in os.listdir(root):
        if os.path.isdir(os.path.join(root, market, name)):
            return True
    return False


def jq_cmd():
    """The image plugin's scripts need jq. Without sudo, the static binary goes into
    ~/.local/bin, which needs no root."""
    if BREW:
        return "brew install jq"
    local = ("mkdir -p ~/.local/bin && curl -fsSL -o ~/.local/bin/jq "
             "https://github.com/jqlang/jq/releases/latest/download/jq-%s-%s "
             "&& chmod +x ~/.local/bin/jq   (~/.local/bin must be on PATH)"
             % ("macos" if sys.platform == "darwin" else "linux",
                "arm64" if platform.machine().lower() in ("arm64", "aarch64") else "amd64"))
    if APT and shutil.which("sudo"):
        return "sudo apt-get install -y jq\n    or, without sudo: " + local
    return local


def image_keys(files=(USER_KEY_FILE, ".env")):
    """{name: (value, source)} for the IMAGE_KEYS set in the environment, then in
    each file in order. Print the names and sources only: this output lands in a
    transcript."""
    found = {k: (os.environ[k], "the environment") for k in IMAGE_KEYS if os.environ.get(k)}
    for path in files:
        if not os.path.isfile(os.path.expanduser(path)):
            continue
        with open(os.path.expanduser(path)) as f:
            for line in f:
                name, eq, value = line.strip().partition("=")
                name = name.replace("export ", "").strip()
                value = value.strip().strip("\"'")
                if eq and name in IMAGE_KEYS and value and name not in found:
                    found[name] = (value, path)
    return found


def check_key(name, value):
    """("valid" | "invalid" | "unchecked", detail). Only an OpenRouter key can be
    asked about itself for free. A key the server could not be asked about is
    unchecked, not invalid."""
    if name != "OPENROUTER_API_KEY":
        return "unchecked", ""
    import json
    import urllib.error
    import urllib.request
    req = urllib.request.Request(OPENROUTER_KEY_URL, headers={"Authorization": "Bearer " + value})
    try:
        urllib.request.urlopen(req, timeout=10).close()
        return "valid", ""
    except urllib.error.HTTPError as e:
        if e.code not in (401, 403):
            return "unchecked", "HTTP %d" % e.code
        try:
            msg = json.loads(e.read().decode("utf-8"))["error"]["message"]
        except (ValueError, KeyError, TypeError):
            msg = e.reason
        return "invalid", ("HTTP %d: %s" % (e.code, msg)).replace(value, "***")
    except (urllib.error.URLError, OSError) as e:
        return "unchecked", str(getattr(e, "reason", e))


def deck_fonts(path):
    from pptx import Presentation
    from collections import Counter
    prs = Presentation(path)
    c = Counter()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.name:
                        c[r.font.name] += len(r.text)
    return [f for f, _n in c.most_common()]


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--font", action="append", default=[],
                    help="a font family the deck needs. Repeat for several.")
    ap.add_argument("--deck", help="read the needed families out of this .pptx instead")
    a = ap.parse_args()

    rows = []  # (ok, required, what, why, fix)

    rows.append((sys.version_info >= (3, 8), True, "python %d.%d" % sys.version_info[:2],
                 "everything here", "install python 3.8 or newer"))

    for mod, pkg, required, why in PY_MODULES:
        rows.append((importlib.util.find_spec(mod) is not None, required,
                     pkg, why, pip_cmd(pkg)))

    soffice = find_soffice()
    rows.append((soffice is not None, True, "libreoffice", "render the deck the way a projector will, which is the only way to see real line breaks",
                 app_cmd("brew install --cask libreoffice", "sudo apt-get install -y libreoffice")))
    if soffice:
        from uno_pdf import find_python
        rows.append((find_python(soffice) is not None, False, "libreoffice uno",
                     "turn off the gaps LibreOffice alone puts between Hangul and Latin text",
                     app_cmd("brew reinstall --cask libreoffice", "sudo apt-get install -y python3-uno")))

    fonts = list(a.font)
    if a.deck and importlib.util.find_spec("pptx"):
        try:
            fonts += [f for f in deck_fonts(a.deck) if f not in fonts]
        except Exception as e:  # a locked or broken file should not stop the report
            print("could not read %s: %s" % (a.deck, e))
    if fonts:
        from fix_orphans import font_file
        for f in fonts:
            rows.append((font_file(f, False) is not None, True, "font %s" % f,
                         "trust a single line break, because a font that is not installed is substituted in silence",
                         "install %s, then copy it into LibreOffice's own font dir:\n"
                         "        /Applications/LibreOffice.app/Contents/Resources/fonts/truetype" % f))

    pdflatex = shutil.which("pdflatex") or (
        "/Library/TeX/texbin/pdflatex" if os.path.exists("/Library/TeX/texbin/pdflatex") else None)
    rows.append((pdflatex is not None, False, "pdflatex",
                 "typeset formulas instead of retyping them in a body font",
                 app_cmd("brew install --cask basictex", "sudo apt-get install -y texlive-latex-base")))

    for name, market, repo, why in SKILL_PLUGINS:
        rows.append((plugin_installed(name), False, "plugin %s" % name, why,
                     "claude plugin marketplace add %s   then   claude plugin install %s@%s"
                     % (repo, name, market)))

    rows.append((shutil.which("jq") is not None, False, "jq",
                 "run the image plugin, whose scripts read every API answer with jq",
                 jq_cmd()))

    # Missing and present-but-invalid are separate rows, and neither prints a value.
    keys = image_keys()
    if not keys:
        rows.append((False, False, "image key",
                     "generate art: the image plugin bills one of these keys",
                     "make one and put it in %s as NAME=..., mode 600, not in\n"
                     "    the project's .env. Never paste it into the chat. Paid, priced per model.\n"
                     % USER_KEY_FILE
                     + "\n".join("        %-19s %s" % (k, url) for k, url in IMAGE_KEYS.items()
                                 if k != "GROK_API_KEY")))
    for name, (value, source) in keys.items():
        state, detail = check_key(name, value)
        rows.append((state != "invalid", False,
                     "image key %s%s, from %s" % (name, {"valid": "", "unchecked": " (unchecked)",
                                                          "invalid": " invalid"}[state], source),
                     "generate art on it. OpenRouter answered %s" % detail,
                     "make a new one at %s,\n    replace it in %s, and never paste it into the chat"
                     % (IMAGE_KEYS[name], source)))

    missing_required = [r for r in rows if not r[0] and r[1]]
    missing_optional = [r for r in rows if not r[0] and not r[1]]

    print("pptx-deck preflight\n")
    for ok, required, what, _why, _fix in rows:
        print("  %s  %-28s %s" % ("ok  " if ok else "MISS", what, "" if required else "(optional)"))

    if missing_required or missing_optional:
        print("")
    for ok, required, what, why, fix in missing_required + missing_optional:
        print("%s %s" % ("REQUIRED" if required else "optional", what))
        print("    without it you cannot %s" % why)
        print("    %s\n" % fix)

    if missing_required:
        print("%d required piece%s missing. Say so before promising a deck."
              % (len(missing_required), "" if len(missing_required) == 1 else "s"))
        return 1
    if missing_optional:
        print("Everything required is here. %d optional capabilit%s unavailable: offer "
              "the install, and run it only on a yes." % (len(missing_optional),
                                               "y is" if len(missing_optional) == 1 else "ies are"))
        return 0
    print("Everything is here.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
