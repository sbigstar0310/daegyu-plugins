#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""A three slide deck in the white reference instance.

Run this first, before you build anything real:

    python3 assets/starter_build.py out.pptx
    python3 scripts/render_real.py out.pptx
    python3 scripts/check_layout.py out.pptx

If all three pass, python-pptx, the font and the renderer work, and you know what
the house grammar looks like. If one fails, fix that before writing a real deck.

This file is deliberately self contained: it imports tokens and python-pptx and
nothing else, so you can copy it anywhere. For a real deck use
`scripts/deck_lib.py` instead of the helpers below, which is the same mechanics
with the edge cases handled.
"""
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import tokens_white as T

HERE = os.path.dirname(os.path.abspath(__file__))
ICONS = os.path.join(HERE, "icons")

PAGE = [0]


# ---------------------------------------------------------------- mechanics
def _highlight(run, color=T.HL):
    """There is no python-pptx API for a marker highlight. a:highlight must come
    BEFORE a:latin in the run properties or PowerPoint drops it silently."""
    rPr = run._r.get_or_add_rPr()
    hl = etree.Element(qn("a:highlight"))
    etree.SubElement(hl, qn("a:srgbClr")).set("val", color)
    latin = rPr.find(qn("a:latin"))
    if latin is not None:
        latin.addprevious(hl)
    else:
        rPr.append(hl)


def tb(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: [{t | runs, size, bold, after, line}]. A run is (text, bold, marker)."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    for i, p_ in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(p_.get("after", T.AFTER))
        p.line_spacing = p_.get("line", T.LINE)
        runs = p_.get("runs") or [(p_["t"], p_.get("bold", False), False)]
        for text, bold, mark in runs:
            if not text:
                continue
            r = p.add_run()
            r.text = text
            r.font.size = Pt(p_.get("size", T.S_BODY))
            r.font.bold = bold
            r.font.color.rgb = T.BLACK
            r.font.name = T.FONT
            if mark:
                _highlight(r)
    return box


def t1(s, x, y, w, h, t, size=T.S_BODY, bold=False, align=PP_ALIGN.LEFT):
    return tb(s, x, y, w, h, [dict(t=t, size=size, bold=bold, after=0)], align=align)


def rect(s, x, y, w, h, fill=None, line=None, lw=0.75):
    """A plain rectangle. Autoshapes inherit a theme fill and a shadow, so the
    p:style element has to go or the deck grows an accent colour it does not have."""
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)
    return shp


def pic(s, path, x, y, w=None, h=None):
    """Returns the placed size, so a caption can be anchored to the real bottom."""
    from PIL import Image
    iw, ih = Image.open(path).size
    if w is None:
        w = h * iw / ih
    if h is None:
        h = w * ih / iw
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


def bullets(s, x, y, w, h, entries, size=T.S_BODY, after=6, indent=0.2):
    """entries: [(lead, text)]. Real bullet paragraphs with a hanging indent, so a
    wrapped second line lines up under the first word and not under the bullet."""
    paras = [dict(runs=[(lead + ": ", True, False), (text, False, False)], size=size, after=after)
             for lead, text in entries]
    box = tb(s, x, y, w, h, paras)
    for p in box.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(indent))))
        pPr.set("indent", str(-int(Inches(indent))))
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
    return box


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt.strip()


# ---------------------------------------------------------------- house grammar
def slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def header(s, title, section=None, number=True):
    """Small grey section label, slide title, page number. Fixed on every slide."""
    if section:
        t1(s, T.ML, T.LABEL_Y, T.CW, 0.2, section, size=T.S_SMALL)
    t1(s, T.ML, T.TITLE_Y, T.CW, 0.44, title, size=T.S_TITLE, bold=True)
    if number:
        PAGE[0] += 1
        t1(s, T.PAGE_X, T.PAGE_Y, T.PAGE_W, 0.18, str(PAGE[0]),
           size=T.S_SMALL, align=PP_ALIGN.RIGHT)


def takeaway(s, text, phrase=None, y=4.76):
    """One line at the bottom, with exactly one phrase on the marker. Optional."""
    if phrase and phrase in text:
        i = text.index(phrase)
        runs = [(text[:i], False, False), (phrase, True, True), (text[i + len(phrase):], False, False)]
    else:
        runs = [(text, True, False)]
    return tb(s, T.ML, y, T.CW, 0.54, [dict(runs=runs, size=T.S_TAKE, after=0)])


def caption(s, x, y, w, text):
    return tb(s, x, y, w, 0.22, [dict(t=text, size=T.S_SMALL, after=0, line=1.0)])


# ---------------------------------------------------------------- the deck
def build(prs):
    # 1. Title. No section label, no page number.
    s = slide(prs)
    t1(s, T.ML, 2.10, T.CW, 0.90, "The white reference instance", size=T.S_SEC, bold=True)
    t1(s, T.ML, 3.10, 6.20, 0.60,
       "A starter deck that proves the toolchain works", size=T.S_BODY)
    rect(s, T.ML, 3.95, T.CW, 0.01, fill=T.GRAY)
    t1(s, T.ML, 4.15, T.CW, 0.24, "pptx-deck", size=T.S_SMALL)
    notes(s, "Three slides. If these render and pass the layout check, "
             "the toolchain is working and you can build the real deck.")

    # 2. Bullets plus icons. Two blocks, so it is inside the density budget.
    s = slide(prs)
    header(s, "Every slide carries a picture", section="Grammar")
    bullets(s, T.ML, 1.50, 5.40, 2.40, [
        ("Header", "a grey label above a bold title, on every slide"),
        ("Body", "at most three blocks, and one of them is a picture"),
        ("Takeaway", "one closing line, one phrase on the marker"),
    ], after=16)
    for i, name in enumerate(["workflow", "shield-check", "gauge"]):
        path = os.path.join(ICONS, name + ".png")
        if os.path.exists(path):
            pic(s, path, 6.55, 1.55 + i * 1.05, w=0.46)
            t1(s, 7.20, 1.60 + i * 1.05, 2.10, 0.30, name, size=T.S_DESC)
    takeaway(s, "A slide with no picture is a slide the audience reads instead of listens to.",
             phrase="reads instead of listens")
    notes(s, "The icons come from assets/icons. They are already ink black at "
             "stroke width 1.6, so they sit at the same weight as the body text.")

    # 3. A card row. Still two blocks: the cards and the takeaway.
    s = slide(prs)
    header(s, "Weight, size, and one marker", section="Grammar")
    cards = [("eye-off", "No accent", "One ink on white. Hierarchy is size and weight."),
             ("scale", "One grey", "#CFCFCF for rules, hairlines and card outlines."),
             ("flag", "One marker", "#E6E6E6 behind a single bold phrase per slide.")]
    cw = (T.CW - 2 * 0.24) / 3
    for i, (icon, head, body) in enumerate(cards):
        x = T.ML + i * (cw + 0.24)
        rect(s, x, 1.60, cw, 2.02, fill=None, line=T.GRAY)
        path = os.path.join(ICONS, icon + ".png")
        if os.path.exists(path):
            pic(s, path, x + 0.22, 1.82, w=0.40)
        t1(s, x + 0.22, 2.40, cw - 0.44, 0.30, head, size=T.S_BODY, bold=True)
        tb(s, x + 0.22, 2.78, cw - 0.44, 0.70, [dict(t=body, size=T.S_DESC, after=0)])
    takeaway(s, "If you reach for a colour, the answer this instance wants is weight.",
             phrase="the answer this instance wants is weight")
    notes(s, "Cards are outlined in the one grey, never filled with colour and "
             "never rounded. Each one still carries an icon.")


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "starter.pptx"
    prs = Presentation()
    prs.slide_width = Inches(T.W)
    prs.slide_height = Inches(T.H)
    build(prs)
    prs.save(out)
    print("wrote %s, %d slides" % (out, len(prs.slides)))


if __name__ == "__main__":
    main()
