# -*- coding: utf-8 -*-
"""The shared helpers of the starter deck: mechanics and house grammar.

Every slide module imports this. It is deck-global, so only the main session
edits it. A slide agent reads it and never changes it: a helper one slide needs
differently is drawn in that slide's own file.

It is deliberately self contained: it imports tokens and python-pptx and nothing
else. For a real deck use `scripts/deck_lib.py` instead of the mechanics below,
which is the same with the edge cases handled, and keep only the grammar
(header, takeaway, caption) in a file like this one.
"""
import os
import sys

from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.dirname(HERE)
sys.path.insert(0, ASSETS)
import tokens_white as T  # noqa: E402

ICONS = os.path.join(ASSETS, "icons")


# ---------------------------------------------------------------- mechanics
def _highlight(run, color=T.HL):
    """There is no python-pptx API for a marker highlight. a:highlight must come
    BEFORE a:latin in the run properties or PowerPoint drops it silently."""
    rPr = run._r.get_or_add_rPr()
    hl = etree.Element(qn("a:highlight"))
    etree.SubElement(hl, qn("a:srgbClr")).set("val", color)
    latin = rPr.find(qn("a:latin"))
    if latin is not None:
        latin.addprevious(hl)
    else:
        rPr.append(hl)


def tb(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: [{t | runs, size, bold, after, line}]. A run is (text, bold, marker)."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    for i, p_ in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(p_.get("after", T.AFTER))
        p.line_spacing = p_.get("line", T.LINE)
        runs = p_.get("runs") or [(p_["t"], p_.get("bold", False), False)]
        for text, bold, mark in runs:
            if not text:
                continue
            r = p.add_run()
            r.text = text
            r.font.size = Pt(p_.get("size", T.S_BODY))
            r.font.bold = bold
            r.font.color.rgb = T.BLACK
            r.font.name = T.FONT
            if mark:
                _highlight(r)
    return box


def t1(s, x, y, w, h, t, size=T.S_BODY, bold=False, align=PP_ALIGN.LEFT):
    return tb(s, x, y, w, h, [dict(t=t, size=size, bold=bold, after=0)], align=align)


def rect(s, x, y, w, h, fill=None, line=None, lw=0.75):
    """A plain rectangle. Autoshapes inherit a theme fill and a shadow, so the
    p:style element has to go or the deck grows an accent colour it does not have."""
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)
    return shp


def pic(s, path, x, y, w=None, h=None):
    """Returns the placed size, so a caption can be anchored to the real bottom."""
    from PIL import Image
    iw, ih = Image.open(path).size
    if w is None:
        w = h * iw / ih
    if h is None:
        h = w * ih / iw
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


def bullets(s, x, y, w, h, entries, size=T.S_BODY, after=6, indent=0.2):
    """entries: [(lead, text)]. Real bullet paragraphs with a hanging indent, so a
    wrapped second line lines up under the first word and not under the bullet."""
    paras = [dict(runs=[(lead + ": ", True, False), (text, False, False)], size=size, after=after)
             for lead, text in entries]
    box = tb(s, x, y, w, h, paras)
    for p in box.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(indent))))
        pPr.set("indent", str(-int(Inches(indent))))
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
    return box


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt.strip()


# ---------------------------------------------------------------- house grammar
def slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def header(s, title, section=None):
    """Small grey section label and slide title. The page number is not drawn
    here: it depends on the slide order, so build_deck.py adds it."""
    if section:
        t1(s, T.ML, T.LABEL_Y, T.CW, 0.2, section, size=T.S_SMALL)
    t1(s, T.ML, T.TITLE_Y, T.CW, 0.44, title, size=T.S_TITLE, bold=True)


def page_number(s, n):
    t1(s, T.PAGE_X, T.PAGE_Y, T.PAGE_W, 0.18, str(n), size=T.S_SMALL, align=PP_ALIGN.RIGHT)


def takeaway(s, text, phrase=None, y=4.76):
    """One line at the bottom, with exactly one phrase on the marker. Optional."""
    if phrase and phrase in text:
        i = text.index(phrase)
        runs = [(text[:i], False, False), (phrase, True, True), (text[i + len(phrase):], False, False)]
    else:
        runs = [(text, True, False)]
    return tb(s, T.ML, y, T.CW, 0.54, [dict(runs=runs, size=T.S_TAKE, after=0)])


def caption(s, x, y, w, text):
    return tb(s, x, y, w, 0.22, [dict(t=text, size=T.S_SMALL, after=0, line=1.0)])
