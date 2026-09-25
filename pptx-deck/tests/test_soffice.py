# -*- coding: utf-8 -*-
"""Tests for where render_real.py and preflight.py look for LibreOffice.

    uv run --directory pptx-deck pytest tests/test_soffice.py

No test finds the soffice installed on the machine, except the one marked
libreoffice. The `nowhere` fixture empties every place the search looks, and a test
puts a stand-in executable in the one place it is about.
"""
import os
import subprocess
import sys

import pytest
from pptx import Presentation
from pptx.util import Inches

import preflight as P
import render_real as R
import soffice as S

SCRIPTS = os.path.dirname(os.path.abspath(S.__file__))


def executable(path):
    """A stand-in soffice at path, which does nothing."""
    path = str(path)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w") as f:
        f.write("#!/bin/sh\nexit 0\n")
    os.chmod(path, 0o755)
    return path


@pytest.fixture
def nowhere(tmp_path, monkeypatch):
    """No variable, an empty PATH, an empty HOME, and no fixed paths or globs. A test
    adds back the one it needs through the returned directories."""
    for var in S.VARIABLES:
        monkeypatch.delenv(var, raising=False)
    path_dir = tmp_path / "bin"
    path_dir.mkdir()
    home = tmp_path / "home"
    home.mkdir()
    monkeypatch.setenv("PATH", str(path_dir))
    monkeypatch.setenv("HOME", str(home))
    monkeypatch.setattr(S, "FIXED", [])
    monkeypatch.setattr(S, "GLOBS", [])
    return {"bin": path_dir, "home": home, "root": tmp_path}


# ---------------------------------------------------------------- the order
def test_the_variable_wins_over_path_and_fixed_paths(nowhere, monkeypatch):
    executable(nowhere["bin"] / "soffice")
    fixed = executable(nowhere["root"] / "fixed" / "soffice")
    monkeypatch.setattr(S, "FIXED", [fixed])
    named = executable(nowhere["root"] / "mine" / "program" / "soffice")
    monkeypatch.setenv("SOFFICE", named)
    assert S.find_soffice() == named


def test_the_namespaced_variable_wins_over_SOFFICE(nowhere, monkeypatch):
    monkeypatch.setenv("SOFFICE", executable(nowhere["root"] / "a" / "soffice"))
    ours = executable(nowhere["root"] / "b" / "soffice")
    monkeypatch.setenv("PPTX_DECK_SOFFICE", ours)
    assert S.find_soffice() == ours


@pytest.mark.parametrize("target", ["missing", "not_executable", "directory"])
def test_a_variable_that_names_no_executable_is_an_error(nowhere, monkeypatch, target):
    # Something the search would find must not stand in for what the user named.
    executable(nowhere["bin"] / "soffice")
    bad = nowhere["root"] / "bad" / "soffice"
    if target == "not_executable":
        bad.parent.mkdir()
        bad.write_text("")
    elif target == "directory":
        bad.mkdir(parents=True)
    monkeypatch.setenv("PPTX_DECK_SOFFICE", str(bad))
    assert S.find_soffice() is None
    msg = S.explain()
    assert "PPTX_DECK_SOFFICE" in msg and str(bad) in msg
    assert "not an executable" in msg


def test_an_empty_variable_is_unset(nowhere, monkeypatch):
    found = executable(nowhere["bin"] / "soffice")
    monkeypatch.setenv("SOFFICE", "")
    assert S.find_soffice() == found


def test_path_wins_over_a_fixed_path(nowhere, monkeypatch):
    # An old distro /usr/bin/soffice must not beat the newer one the user put first.
    monkeypatch.setattr(S, "FIXED", [executable(nowhere["root"] / "usr" / "bin" / "soffice")])
    on_path = executable(nowhere["bin"] / "soffice")
    assert S.find_soffice() == on_path


def test_libreoffice_on_path_is_found_too(nowhere):
    assert S.find_soffice() is None
    found = executable(nowhere["bin"] / "libreoffice")
    assert S.find_soffice() == found


def test_a_fixed_path_is_used_when_path_has_none(nowhere, monkeypatch):
    absent = str(nowhere["root"] / "Applications" / "soffice")
    fixed = executable(nowhere["root"] / "usr" / "local" / "bin" / "soffice")
    monkeypatch.setattr(S, "FIXED", [absent, fixed])
    assert S.find_soffice() == fixed


def test_a_fixed_path_under_home_is_expanded(nowhere, monkeypatch):
    app = executable(nowhere["home"] / "Applications" / "LibreOffice.app" / "Contents"
                     / "MacOS" / "soffice")
    monkeypatch.setattr(S, "FIXED", ["~/Applications/LibreOffice.app/Contents/MacOS/soffice"])
    assert S.find_soffice() == app


def test_opt_glob_finds_the_official_package(nowhere, monkeypatch):
    opt = nowhere["root"] / "opt"
    found = executable(opt / "libreoffice26.2" / "program" / "soffice")
    monkeypatch.setattr(S, "GLOBS", [str(opt / "libreoffice*" / "program" / "soffice")])
    assert S.find_soffice() == found


def test_the_newest_version_wins_numerically(nowhere, monkeypatch):
    # As strings "7.3" sorts after "26.2", and "7.10" before "7.9".
    opt = nowhere["root"] / "opt"
    for v in ("7.3", "7.10", "26.2", "7.9"):
        executable(opt / ("libreoffice" + v) / "program" / "soffice")
    monkeypatch.setattr(S, "GLOBS", [str(opt / "libreoffice*" / "program" / "soffice")])
    assert S.find_soffice() == str(opt / "libreoffice26.2" / "program" / "soffice")


def test_a_user_local_install_is_found(nowhere, monkeypatch):
    # An extracted .deb, with no root: the version sits deep inside the tree.
    found = executable(nowhere["home"] / ".local" / "opt" / "libreoffice" / "opt"
                       / "libreoffice26.2" / "program" / "soffice")
    monkeypatch.setattr(S, "GLOBS", ["~/.local/opt/libreoffice*/**/program/soffice"])
    assert S.find_soffice() == found


def test_the_default_places_cover_the_issue(nowhere):
    assert "~/Applications/LibreOffice.app/Contents/MacOS/soffice" in S.DEFAULT_FIXED
    assert "/usr/bin/soffice" in S.DEFAULT_FIXED
    assert "/opt/libreoffice*/program/soffice" in S.DEFAULT_GLOBS
    assert "~/.local/opt/libreoffice*/**/program/soffice" in S.DEFAULT_GLOBS


# ---------------------------------------------------------------- the message
def test_the_error_lists_where_it_looked_and_names_the_variable(nowhere, monkeypatch):
    monkeypatch.setattr(S, "FIXED", ["/nowhere/usr/bin/soffice"])
    monkeypatch.setattr(S, "GLOBS", ["/nowhere/opt/libreoffice*/program/soffice"])
    assert S.find_soffice() is None
    msg = S.explain()
    for place in ("$PPTX_DECK_SOFFICE", "$SOFFICE", "PATH", "/nowhere/usr/bin/soffice",
                  "/nowhere/opt/libreoffice*/program/soffice"):
        assert place in msg
    assert "export PPTX_DECK_SOFFICE=/path/to/program/soffice" in msg
    assert "not installed" not in msg


def test_explain_is_none_when_it_is_found(nowhere):
    executable(nowhere["bin"] / "soffice")
    assert S.explain() is None


def test_render_real_says_where_it_looked(tmp_path, nowhere, monkeypatch, capsys):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(tmp_path / "d.pptx"))
    monkeypatch.setattr(sys, "argv", ["render_real.py", str(tmp_path / "d.pptx"),
                                      "-o", str(tmp_path / "render")])
    assert R.main() == 2
    err = capsys.readouterr().err
    assert "export PPTX_DECK_SOFFICE" in err
    assert "not installed" not in err


def test_preflight_shows_the_hint_when_it_is_missing(nowhere, monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["preflight.py"])
    P.main()
    out = capsys.readouterr().out
    assert "MISS  libreoffice" in out
    assert "export PPTX_DECK_SOFFICE" in out


def test_preflight_shows_the_soffice_it_found(nowhere, monkeypatch, capsys):
    found = executable(nowhere["bin"] / "soffice")
    monkeypatch.setattr(sys, "argv", ["preflight.py"])
    P.main()
    line = next(l for l in capsys.readouterr().out.splitlines() if "libreoffice" in l)
    assert line.startswith("  ok") and found in line


# ---------------------------------------------------------------- one finder
def test_render_real_and_preflight_use_the_same_finder():
    assert R.find_soffice is S.find_soffice
    assert P.find_soffice is S.find_soffice


def test_soffice_and_preflight_import_without_the_libraries_preflight_checks():
    # preflight exists to say that PIL or python-pptx is missing, so it must import
    # without them, and so must the finder it shares with render_real.
    code = ("import sys\n"
            "for m in ('PIL', 'pptx', 'lxml', 'pymupdf', 'fitz', 'fix_orphans'):\n"
            "    sys.modules[m] = None\n"
            "sys.path.insert(0, %r)\n"
            "import soffice, preflight\n"
            "print('ok')\n" % SCRIPTS)
    r = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    assert r.stdout.strip() == "ok"


def test_render_real_prints_the_soffice_it_used(tmp_path, monkeypatch, capsys):
    import test_render_real as T
    path = T.deck(tmp_path / "d.pptx", hidden=())
    calls = []

    def run(cmd, capture_output=False, text=False):
        calls.append(cmd)
        import pymupdf
        doc = pymupdf.open()
        doc.new_page(width=200, height=100)
        for _ in range(4):
            doc.new_page(width=200, height=100)
        doc.save(os.path.join(cmd[5], "d.pdf"))
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(R, "find_soffice", lambda: "/opt/lo/program/soffice")
    # The plain export, so the command line starts with the soffice that was found.
    monkeypatch.setattr(R, "find_uno_python", lambda soffice: None)
    monkeypatch.setattr(R.subprocess, "run", run)
    monkeypatch.setattr(sys, "argv", ["render_real.py", path, "-o", str(tmp_path / "r"),
                                      "--dpi", "72"])
    assert R.main() == 0
    assert "soffice /opt/lo/program/soffice" in capsys.readouterr().out.splitlines()
    assert calls[0][0] == "/opt/lo/program/soffice"


# ---------------------------------------------------------------- real LibreOffice
REAL = S.find_soffice()


@pytest.mark.libreoffice
@pytest.mark.skipif(REAL is None, reason="LibreOffice is not installed")
def test_a_real_render_through_the_variable(tmp_path, nowhere, monkeypatch, capsys):
    # PATH, the fixed paths and the globs are empty: only the variable can find it.
    monkeypatch.setenv("PPTX_DECK_SOFFICE", REAL)
    prs = Presentation()
    tb = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(6), Inches(1))
    tb.text_frame.text = "THROUGH THE VARIABLE"
    prs.save(str(tmp_path / "d.pptx"))
    monkeypatch.setattr(sys, "argv", ["render_real.py", str(tmp_path / "d.pptx"),
                                      "-o", str(tmp_path / "render"), "--dpi", "72",
                                      "--keep-pdf"])
    assert R.main() == 0
    assert "soffice %s" % REAL in capsys.readouterr().out.splitlines()
    import pymupdf
    doc = pymupdf.open(str(tmp_path / "render" / "d.pdf"))
    assert doc[0].get_text().strip() == "THROUGH THE VARIABLE"
    assert (tmp_path / "render" / "s01.png").exists()
