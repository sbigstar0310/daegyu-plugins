# -*- coding: utf-8 -*-
"""Tests for the language checks in check_layout.py.

    uv run --directory pptx-deck pytest tests/test_language.py

The character lint runs by default, because it enforces a standing rule. The
sentence-length check runs with --lang-check. The two characters are built with
chr() here: check_layout.py is the only file in the skill that spells them out.
"""
import sys

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

import check_layout as C

DASH = chr(0x2014)   # em dash
DOT = chr(0xB7)      # middle dot


def words(n, first="word"):
    """A sentence of exactly n words, ending in a period."""
    return " ".join([first] + ["w%d" % k for k in range(2, n + 1)]) + "."


def deck(*slides):
    """One slide per argument. A slide is (paragraphs, notes): paragraphs is a list
    of strings for one text box, notes a string or None."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    for paras, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tf = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3)).text_frame
        for i, text in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = text
            r.font.size = Pt(18)
        if notes is not None:
            slide.notes_slide.notes_text_frame.text = notes
    return prs


def chars(prs):
    rep = C.Report()
    C.check_language(prs, rep)
    return [r[3] for r in rep.rows]


def sentences(prs, max_words=None):
    rep = C.Report()
    C.check_sentences(prs, rep, C.MAX_WORDS if max_words is None else max_words)
    return [(r[0], r[2], r[3]) for r in rep.rows]


@pytest.fixture
def run_main(fonts, tmp_path, monkeypatch, capsys):
    def run(prs, *flags):
        path = str(tmp_path / "deck.pptx")
        prs.save(path)
        monkeypatch.setattr(sys, "argv", ["check_layout.py", path,
                                          "--margins", "0.17,9.83,5.31"] + list(flags))
        code = C.main()
        return code, capsys.readouterr().out
    return run


def language_rows(out):
    return [line for line in out.splitlines() if line.split()[1:2] == ["language"]]


# ---------------------------------------------------------------- characters
def test_the_character_lint_runs_without_the_flag(run_main):
    prs = deck((["Results %s runner" % DASH], None))
    _code, out = run_main(prs)
    assert len(language_rows(out)) == 1
    assert "em dash in: Results %s runner" % DASH in out


def test_only_can_leave_the_character_lint_out(run_main):
    prs = deck((["Results %s runner" % DASH], None))
    _code, out = run_main(prs, "--only", "bounds")
    assert language_rows(out) == []


def test_a_character_fails_strict(run_main):
    prs = deck((["runner %s verifier" % DOT], None))
    code, _out = run_main(prs, "--strict")
    assert code == 1


def test_each_paragraph_with_the_character_is_reported():
    prs = deck((["first %s one" % DASH, "clean", "second %s two" % DASH], None))
    assert chars(prs) == ["em dash in: first %s one" % DASH,
                          "em dash in: second %s two" % DASH]


def test_both_characters_in_one_paragraph_are_both_named():
    prs = deck((["Results %s runner %s verifier" % (DASH, DOT)], None))
    assert [m.split(" in: ")[0] for m in chars(prs)] == ["em dash", "middle dot separator"]


def test_notes_are_linted_too():
    prs = deck((["clean"], "say it %s slowly" % DASH))
    assert chars(prs) == ["em dash in the notes: say it %s slowly" % DASH]


# ---------------------------------------------------------------- sentence length
def test_a_long_sentence_is_reported_with_its_length():
    prs = deck(([words(36, "The")], None))
    [(level, slide, msg)] = sentences(prs)
    assert (level, slide) == ("WARN", 1)
    assert msg.startswith("longest sentence is 36 words (limit 20): The w2 w3")


def test_only_the_longest_sentence_per_slide_is_reported():
    prs = deck(([words(25, "a"), words(30, "b") + " " + words(22, "c")], None))
    assert [m for _l, _s, m in sentences(prs)] == [
        "longest sentence is 30 words (limit 20): " + words(30, "b")[:48] + "..."]


def test_a_sentence_at_the_limit_is_clean():
    prs = deck(([words(12), words(20)], words(20)))
    assert sentences(prs) == []


def test_notes_are_measured_separately():
    prs = deck(([words(12)], words(24, "Say")),
               ([words(21, "Show")], words(5)))
    assert [(s, m[:44]) for _l, s, m in sentences(prs)] == [
        (1, "longest sentence in the notes is 24 words (l"),
        (2, "longest sentence is 21 words (limit 20): Sho")]


def test_each_bullet_is_its_own_sentence():
    bullet = " ".join("b%d" % k for k in range(15))      # 15 words, no period
    prs = deck(([bullet, bullet, bullet], None))
    assert sentences(prs) == []


def test_question_and_exclamation_marks_end_a_sentence():
    prs = deck(([words(15)[:-1] + "? " + words(15)[:-1] + "! " + words(15)], None))
    assert sentences(prs) == []


def test_korean_sentences_are_split_and_counted():
    one = " ".join(["결과를"] * 12) + " 비교했다."       # 13 words
    two = " ".join(["검증기가"] * 21) + " 확인한다."     # 22 words
    prs = deck(([one + " " + two], None))
    [(_l, _s, msg)] = sentences(prs)
    assert msg.startswith("longest sentence is 22 words (limit 20): 검증기가")


def test_max_words_sets_the_limit():
    prs = deck(([words(36)], None))
    assert sentences(prs, 40) == []
    assert len(sentences(prs, 35)) == 1


def test_sentence_length_needs_the_flag(run_main):
    prs = deck(([words(36)], None))
    _code, out = run_main(prs)
    assert language_rows(out) == []
    _code, out = run_main(prs, "--lang-check")
    assert "longest sentence is 36 words (limit 20)" in out
    _code, out = run_main(prs, "--lang-check", "--max-words", "40")
    assert language_rows(out) == []
