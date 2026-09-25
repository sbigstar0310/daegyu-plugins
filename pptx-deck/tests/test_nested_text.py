# -*- coding: utf-8 -*-
"""Tests for text inside groups and tables (#11): where each frame lands on the
slide, that every text check reaches it, and how far a table grows.

    uv run --directory pptx-deck pytest tests/test_nested_text.py

The test fonts make each character half an em wide, so at 12 pt a character and a
space are 6 pt each. A 4 in column with 0.1 in side margins holds 273.6 pt, 45
characters, per line. A cell's top and bottom margins are 0.05 in, and a 12 pt
line at 1.0 spacing is 14.4 pt, 0.2 in.
"""
import sys

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

import check_layout as C
import fix_orphans as M

EM_DASH = chr(0x2014)
MIDDLE_DOT = chr(0xB7)
# 44 characters fill the first line to 264 pt, and "yy" is left alone on the
# second: 12 pt, under a third of 273.6. Still two lines at plus 3 percent.
ORPHANED = "x" * 44 + " yy"


def deck(slides=2):
    """A deck whose first slide is blank: the orphan check skips slide 1."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    for _ in range(slides):
        prs.slides.add_slide(prs.slide_layouts[6])
    return prs


def fill(tf, lines, size=12, family="Inter"):
    """One paragraph per line, single spaced, no space after."""
    for i, text in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.0
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.name = family
    return tf


def text_box(shapes, left, top, width, height, lines, name="box", **kw):
    tb = shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.1))
    fill(tf, lines, **kw)
    return tb


def place(group, off, ext, ch_off, ch_ext):
    """Set a group's transform, in inches. python-pptx recomputes it whenever a
    child is added, so this comes after the children."""
    xfrm = group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    for tag, (a, b), names in (("a:off", off, ("x", "y")), ("a:ext", ext, ("cx", "cy")),
                               ("a:chOff", ch_off, ("x", "y")), ("a:chExt", ch_ext, ("cx", "cy"))):
        el = xfrm.find(qn(tag))
        el.set(names[0], str(Inches(a)))
        el.set(names[1], str(Inches(b)))


def table(slide, rows, cols, top=1.0, width=4.0, row_h=0.30, name="Table"):
    gf = slide.shapes.add_table(rows, cols, Inches(1), Inches(top), Inches(width),
                                Inches(row_h * rows))
    gf.name = name
    return gf


def cell_text(gf, r, c, lines, **kw):
    tf = gf.table.cell(r, c).text_frame
    fill(tf, lines, **kw)
    return tf


def inches(frame):
    return tuple(round(v / 914400.0, 4) for v in (frame.left, frame.top, frame.width, frame.height))


# ---------------------------------------------------------------- where frames land
def test_a_group_child_is_placed_on_the_slide():
    prs = deck(1)
    grp = prs.slides[0].shapes.add_group_shape()
    text_box(grp.shapes, 2, 2, 2, 1, ["child"], name="Child")
    place(grp, (1, 1), (4, 2), (0, 0), (8, 4))
    (frame,) = M.text_frames(prs.slides[0].shapes)
    assert frame.name == "Child"
    assert inches(frame) == (2.0, 2.0, 1.0, 0.5)


def test_nested_groups_compose():
    prs = deck(1)
    outer = prs.slides[0].shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    text_box(inner.shapes, 2, 2, 4, 2, ["deep"], name="Deep")
    # Inner: child (2, 2) w 4 -> 2 + 2 x 0.5 = 3 in the outer group, w 2.
    place(inner, (2, 2), (2, 2), (0, 0), (4, 4))
    # Outer: 3 -> 0 + 3 x 0.5 = 1.5 on the slide, w 1.
    place(outer, (0, 0), (4, 4), (0, 0), (8, 8))
    (frame,) = M.text_frames(prs.slides[0].shapes)
    assert inches(frame) == (1.5, 1.5, 1.0, 0.5)


def test_a_cell_is_its_column_minus_its_margins():
    prs = deck(1)
    gf = table(prs.slides[0], 2, 2, width=8.0)
    frames = list(M.text_frames(prs.slides[0].shapes))
    assert [f.key[1:] for f in frames] == [(0, 0), (0, 1), (1, 0), (1, 1)]
    first = frames[0]
    assert first.kind == "cell" and first.wrap
    assert inches(first) == (1.0, 1.0, 4.0, 0.3)
    cell_text(gf, 0, 0, ["one"])
    p = first.text_frame.paragraphs[0]
    assert M.para_width(first, p) == pytest.approx(273.6)


def test_a_merged_cell_spans_its_columns_and_the_hidden_ones_are_skipped():
    prs = deck(1)
    gf = table(prs.slides[0], 2, 3)
    for k, w in enumerate((2, 3, 4)):
        gf.table.columns[k].width = Inches(w)
    gf.table.cell(0, 0).merge(gf.table.cell(0, 1))
    frames = list(M.text_frames(prs.slides[0].shapes))
    assert [f.key[1:] for f in frames] == [(0, 0), (0, 2), (1, 0), (1, 1), (1, 2)]
    assert inches(frames[0])[2] == 5.0
    assert inches(frames[1])[:3] == (6.0, 1.0, 4.0)


# ---------------------------------------------------------------- overflow
def overflow(prs):
    rep = C.Report()
    C.check_overflow(prs, rep, "Inter")
    return rep


# Four 20-character words: one per line in a 1.8 in inner width (129.6 pt), so
# 0.2 in of insets and four 0.2 in lines need 1.00 in.
CAPTION = " ".join(c * 20 for c in "abcd")


def test_a_grouped_caption_overflows_like_an_ungrouped_one(fonts):
    fonts("Inter.ttf", "Inter", "Regular")
    prs = deck(2)
    grp = prs.slides[0].shapes.add_group_shape()
    text_box(grp.shapes, 1, 1, 2.0, 0.30, [CAPTION], name="Grouped_caption")
    grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.5), Inches(1), Inches(1), Inches(1))
    text_box(prs.slides[1].shapes, 1, 1, 2.0, 0.30, [CAPTION], name="Top_level_text")
    assert [r[2:] for r in overflow(prs).errors()] == [
        (1, "Grouped_caption needs 1.00 in of height and has 0.30"),
        (2, "Top_level_text needs 1.00 in of height and has 0.30")]


def test_a_scaled_group_wraps_at_its_drawn_width(fonts):
    """8 in wide in the group, drawn 4 in wide: 360 pt of text is two lines at the
    drawn width and one at the child's own."""
    fonts("Inter.ttf", "Inter", "Regular")
    prs = deck(1)
    grp = prs.slides[0].shapes.add_group_shape()
    text_box(grp.shapes, 0, 0, 8, 0.8, ["x" * 29 + " " + "y" * 30], name="Scaled")
    place(grp, (1, 1), (4, 0.4), (0, 0), (8, 0.8))
    assert [r[3] for r in overflow(prs).errors()] == [
        "Scaled needs 0.60 in of height and has 0.40"]


# ---------------------------------------------------------------- orphans and --fix
def test_an_orphan_in_a_cell_is_reported(fonts):
    fonts("Inter.ttf", "Inter", "Regular")
    prs = deck()
    gf = table(prs.slides[1], 1, 1, row_h=0.6)
    cell_text(gf, 0, 0, [ORPHANED])
    assert M.report(prs) == [(2, ORPHANED, "orphan")]


def test_fix_shrinks_the_font_in_a_cell_and_never_the_column(fonts):
    """At 11.5 pt the text is 270.25 pt and fits one line of 273.6."""
    fonts("Inter.ttf", "Inter", "Regular")
    prs = deck()
    gf = table(prs.slides[1], 1, 1, row_h=0.6)
    tf = cell_text(gf, 0, 0, [ORPHANED])
    assert M.fix(prs) == 1
    assert tf.paragraphs[0].runs[0].font.size == Pt(11.5)
    assert gf.table.columns[0].width == Inches(4)
    assert gf.width == Inches(4)
    # With the font held, pass 2 would narrow the box, and a cell has none of its own.
    tf.paragraphs[0].runs[0].font.size = Pt(12)
    assert M.fix(prs, min_font_scale=1.0) == 0
    assert gf.table.columns[0].width == Inches(4)


def test_an_orphan_in_a_group_is_reported_and_fixed(fonts):
    fonts("Inter.ttf", "Inter", "Regular")
    prs = deck()
    grp = prs.slides[1].shapes.add_group_shape()
    tb = text_box(grp.shapes, 1, 1, 4, 0.6, [ORPHANED], name="Grouped")
    grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6), Inches(1), Inches(1), Inches(1))
    assert M.report(prs) == [(2, ORPHANED, "orphan")]
    assert M.fix(prs) == 1
    assert tb.text_frame.paragraphs[0].runs[0].font.size == Pt(11.5)
    # One line now, 99 percent full: at risk, as the same frame at the top level is.
    assert [k for _n, _t, k in M.report(prs)] == ["at risk"]


def test_a_group_child_is_narrowed_in_its_own_coordinates(fonts):
    """Five 10-character words in a child 8 in wide, drawn 4 in wide: four words
    (258 pt) on the first line and one left over. With the font held, the drawn
    frame narrows until three words fill a line, at 0.90: 244.8 pt, and 252.1 at
    plus 3 percent. The child keeps 0.90 of its own 8 in, so it is drawn 3.6 in."""
    fonts("Inter.ttf", "Inter", "Regular")
    prs = deck()
    grp = prs.slides[1].shapes.add_group_shape()
    tb = text_box(grp.shapes, 0, 0, 8, 1.2, [" ".join(c * 10 for c in "abcde")], name="Wide")
    place(grp, (1, 1), (4, 0.6), (0, 0), (8, 1.2))
    assert [k for _n, _t, k in M.report(prs)] == ["orphan"]
    assert M.fix(prs, min_font_scale=1.0) == 1
    assert tb.width == pytest.approx(Inches(7.2), abs=10)
    (frame,) = M.text_frames(prs.slides[1].shapes)
    assert frame.width == pytest.approx(Inches(3.6), abs=10)
    assert M.report(prs) == []


# ---------------------------------------------------------------- type, language
def test_a_small_run_in_a_cell_or_group_is_below_the_floor():
    prs = deck(1)
    slide = prs.slides[0]
    gf = table(slide, 1, 1, row_h=0.6)
    cell_text(gf, 0, 0, ["tiny cell"], size=6)
    grp = slide.shapes.add_group_shape()
    text_box(grp.shapes, 6, 1, 2, 0.5, ["tiny group"], size=6)
    grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6), Inches(2), Inches(1), Inches(1))
    rep = C.Report()
    sizes = C.check_type(prs, rep, C.MIN_PT)
    assert sizes == {6.0: len("tiny cell") + len("tiny group")}
    assert rep.warnings() == [
        ("WARN", "type", 1, "2 runs at 6.0 pt, below the floor of 8.0, such as 'tiny cell'")]


def test_an_em_dash_in_a_cell_and_a_middle_dot_in_a_group():
    prs = deck(1)
    slide = prs.slides[0]
    gf = table(slide, 1, 1, row_h=0.6)
    cell_text(gf, 0, 0, ["field by field %s twice" % EM_DASH])
    grp = slide.shapes.add_group_shape()
    text_box(grp.shapes, 6, 1, 2, 0.5, ["runner %s verifier" % MIDDLE_DOT])
    grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6), Inches(2), Inches(1), Inches(1))
    rep = C.Report()
    C.check_language(prs, rep)
    assert [r[3] for r in rep.warnings()] == [
        "em dash in: field by field %s twice" % EM_DASH,
        "middle dot separator in: runner %s verifier" % MIDDLE_DOT]


# ---------------------------------------------------------------- render_real
def test_render_real_sees_a_font_named_only_in_a_cell(tmp_path, fonts, monkeypatch, capsys):
    import render_real as R
    prs = deck(1)
    gf = table(prs.slides[0], 1, 1, row_h=0.6)
    cell_text(gf, 0, 0, ["drawn in something else"], family="Comic Neue")
    assert R.families(prs) == [("Comic Neue", False)]
    path = str(tmp_path / "cell_font.pptx")
    prs.save(path)
    monkeypatch.setattr(R, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(sys, "argv", ["render_real.py", path, "-o", str(tmp_path / "r")])
    assert R.main() == 2
    assert "Comic Neue" in capsys.readouterr().err


# ---------------------------------------------------------------- tables that grow
THREE = ["line one", "line two", "line three"]


def grown_deck(lines, top=1.0):
    """A 1 x 1 table, 4 in wide, one 0.30 in row from `top`."""
    prs = deck(1)
    gf = table(prs.slides[0], 1, 1, top=top)
    cell_text(gf, 0, 0, lines)
    return prs, gf


def test_a_row_that_grows_is_reported(fonts):
    """0.05 + 3 x 0.2 + 0.05 = 0.70 in, where 0.30 was declared."""
    fonts("Inter.ttf", "Inter", "Regular")
    prs, _gf = grown_deck(THREE)
    rep = overflow(prs)
    assert rep.errors() == []
    assert rep.warnings() == [
        ("WARN", "overflow", 1,
         "Table row 1 grows from 0.30 to 0.70 in; table bottom 1.30 to 1.70 in")]


def test_a_row_that_fits_is_not(fonts):
    fonts("Inter.ttf", "Inter", "Regular")
    prs, _gf = grown_deck(["line one"])
    assert overflow(prs).rows == []


def test_the_grown_table_crosses_the_content_line(fonts):
    fonts("Inter.ttf", "Inter", "Regular")
    prs, _gf = grown_deck(THREE)
    rep = C.Report()
    C.check_bounds(prs, rep, 0.5, 9.5, 1.5, "Inter")
    assert rep.warnings() == [
        ("WARN", "bounds", 1, "Table ends at 1.70, below the content line 1.50")]
    # The declared box alone ends at 1.30, above the line.
    prs, _gf = grown_deck(["line one"])
    rep = C.Report()
    C.check_bounds(prs, rep, 0.5, 9.5, 1.5, "Inter")
    assert rep.rows == []


def test_the_grown_table_overlaps_the_shape_below(fonts):
    fonts("Inter.ttf", "Inter", "Regular")
    for lines, expected in ((["line one"], []),
                            (THREE, ["Table and Below overlap by 4.00 x 0.25 in"])):
        prs, _gf = grown_deck(lines)
        below = prs.slides[0].shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.45),
                                               Inches(4), Inches(0.55))
        below.name = "Below"
        below.fill.solid()
        below.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
        rep = C.Report()
        C.check_overlap(prs, rep, C.EPS, "Inter")
        assert [r[3] for r in rep.errors()] == expected


# ---------------------------------------------------------------- real LibreOffice
@pytest.mark.libreoffice
def test_a_grown_table_renders_where_it_was_predicted(tmp_path):
    """A 0.30 in row holding five lines of Inter. LibreOffice grows it, and the
    bottom of the drawn table must sit within 0.1 in of the predicted one."""
    import subprocess

    import pymupdf
    import render_real as R

    soffice = R.find_soffice()
    if soffice is None:
        pytest.skip("LibreOffice is not installed")
    M._file_cache.clear()
    M._name_table = None
    if M.font_file("Inter", False) is None:
        pytest.skip("Inter is not installed, so LibreOffice would draw another font")

    prs, gf = grown_deck(["row line %d" % k for k in range(1, 6)])
    path = tmp_path / "grown.pptx"
    prs.save(str(path))
    rep = overflow(prs)
    (row,) = [r[3] for r in rep.warnings() if "grows" in r[3]]
    predicted = float(row.rsplit(" to ", 1)[1].split()[0])

    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir",
                    str(tmp_path), str(path)], capture_output=True, check=True)
    page = pymupdf.open(str(tmp_path / "grown.pdf"))[0]
    assert any("Inter" in f[3] for f in page.get_fonts()), page.get_fonts()
    per_inch = page.rect.width / 10.0
    # The cell fill, the only thing drawn from the table's left edge at 1 in.
    bottoms = [d["rect"].y1 for d in page.get_drawings()
               if d.get("fill") is not None and abs(d["rect"].x0 / per_inch - 1.0) < 0.05]
    assert bottoms, "LibreOffice drew no table fill"
    drawn = max(bottoms) / per_inch
    assert drawn == pytest.approx(predicted, abs=0.1)
    assert drawn > 1.9                     # it did grow: 1.30 was declared
