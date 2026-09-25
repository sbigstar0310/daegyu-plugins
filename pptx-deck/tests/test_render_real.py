# -*- coding: utf-8 -*-
"""Tests for how render_real.py treats hidden slides.

    uv run --directory pptx-deck pytest
    uv run --directory pptx-deck pytest -m "not libreoffice"    # skip the real render

Most tests stand in for LibreOffice: the fake writes one PDF page per exported
slide, 200 + 10 x slide number points wide, so the width of a PNG rendered at 72
dpi says which slide is on it. A PNG named s03.png must be 230 px wide.
"""
import os
import re
import subprocess
import sys

import pymupdf
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import render_real as R

HIDDEN = (2, 5)


def deck(path, hidden=HIDDEN, n=5, show="0"):
    prs = Presentation()
    for i in range(1, n + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        tb.text_frame.text = "SLIDE %d%s" % (i, " HIDDEN" if i in hidden else "")
        if i in hidden:
            slide._element.set("show", show)
    prs.save(str(path))
    return str(path)


@pytest.fixture
def libreoffice(monkeypatch):
    """A stand-in for soffice. It decides what is hidden the way LibreOffice does,
    not by asking render_real. honours_filter=False plays a LibreOffice older than
    7.4, which leaves hidden slides out regardless; max_pages cuts the PDF short;
    writes=False plays a LibreOffice that fails to write and still exits 0."""
    calls = []
    state = {"honours_filter": True, "max_pages": None, "writes": True}

    def run(cmd, capture_output=False, text=False):
        calls.append(cmd)
        if not state["writes"]:
            return subprocess.CompletedProcess(cmd, 0, "", "Error: source file could not be loaded")
        target, outdir, path = cmd[3], cmd[5], cmd[6]
        with_hidden = "ExportHiddenSlides" in target and state["honours_filter"]
        doc = pymupdf.open()
        for i, slide in enumerate(Presentation(path).slides):
            if slide._element.get("show") in ("0", "false") and not with_hidden:
                continue
            if state["max_pages"] is not None and doc.page_count == state["max_pages"]:
                break
            doc.new_page(width=200 + 10 * (i + 1), height=100)
        doc.save(os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".pdf"))
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(R, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(R.subprocess, "run", run)

    def configure(**kw):
        state.update(kw)
        return calls

    return configure


@pytest.fixture
def render(tmp_path, monkeypatch, capsys):
    out = tmp_path / "render"

    def run(path, *flags):
        monkeypatch.setattr(sys, "argv", ["render_real.py", path, "-o", str(out),
                                          "--dpi", "72"] + list(flags))
        code = R.main()
        io = capsys.readouterr()
        # A run stopped before it starts, as the font gate does, makes no directory.
        names = sorted(os.listdir(str(out))) if out.exists() else []
        pngs = {name: Image.open(str(out / name)).size[0]
                for name in names if name.endswith(".png")}
        return code, io.out, io.err, pngs

    return run


def slide_width(n):
    return 200 + 10 * n


# ---------------------------------------------------------------- the map
@pytest.mark.parametrize("skip_hidden, expected", [
    pytest.param(False, [(1, 1), (2, 2), (3, 3), (4, 4), (5, 5)], id="rendered"),
    pytest.param(True, [(1, 1), (2, None), (3, 2), (4, 3), (5, None)], id="skipped"),
])
def test_slide_map(tmp_path, skip_hidden, expected):
    prs = Presentation(deck(tmp_path / "d.pptx"))
    assert R.slide_map(prs, skip_hidden) == expected


# ---------------------------------------------------------------- the command
def test_hidden_slides_are_rendered_by_default(tmp_path, libreoffice, render):
    calls = libreoffice()
    code, out, err, pngs = render(deck(tmp_path / "d.pptx"))
    assert code == 0
    assert calls[0][3] == R.PDF_WITH_HIDDEN
    assert pngs == {"s%02d.png" % n: slide_width(n) for n in range(1, 6)}
    assert "hidden  2, 5, rendered like the rest, so page N is slide N." in out
    assert err == ""


def test_skip_hidden_renders_what_the_slideshow_shows(tmp_path, libreoffice, render):
    calls = libreoffice()
    code, out, err, pngs = render(deck(tmp_path / "d.pptx"), "--skip-hidden")
    assert code == 0
    assert calls[0][3] == "pdf"
    assert pngs == {"s%02d.png" % n: slide_width(n) for n in (1, 3, 4)}
    assert "slide to page map" in out
    assert "1>1  2>hidden  3>2  4>3  5>hidden" in out
    assert err == ""


def test_a_libreoffice_that_ignores_the_option_does_not_misname_pages(tmp_path, libreoffice,
                                                                      render):
    libreoffice(honours_filter=False)
    code, out, err, pngs = render(deck(tmp_path / "d.pptx"))
    assert code == 0
    assert "LibreOffice 7.4 or newer" in err
    assert "guess" not in err
    assert pngs == {"s%02d.png" % n: slide_width(n) for n in (1, 3, 4)}
    assert "2>hidden" in out


def test_slides_are_picked_by_slide_number(tmp_path, libreoffice, render):
    libreoffice()
    _code, _out, _err, pngs = render(deck(tmp_path / "d.pptx"), "--slides", "5")
    assert pngs == {"s05.png": slide_width(5)}
    # Skipping hidden slides, slide 4 is on page 3.
    _code, _out, _err, pngs = render(deck(tmp_path / "d.pptx"), "--skip-hidden", "--slides", "4")
    assert pngs == {"s04.png": slide_width(4)}


def test_a_slide_hidden_with_false_is_hidden(tmp_path, libreoffice, render):
    libreoffice()
    code, out, err, pngs = render(deck(tmp_path / "d.pptx", show="false"), "--skip-hidden")
    assert code == 0
    assert pngs == {"s%02d.png" % n: slide_width(n) for n in (1, 3, 4)}
    assert err == ""


def test_an_old_pdf_is_not_rendered_in_place_of_a_failed_export(tmp_path, libreoffice, render):
    libreoffice()
    path = deck(tmp_path / "d.pptx")
    render(path, "--skip-hidden", "--keep-pdf")
    assert (tmp_path / "render" / "d.pdf").exists()
    libreoffice(writes=False)
    code, _out, err, pngs = render(path)
    assert code == 2
    assert "LibreOffice produced no PDF" in err
    assert "7.4" not in err
    assert pngs == {}


def test_a_short_pdf_does_not_crash_the_guess(tmp_path, libreoffice, render):
    libreoffice(max_pages=2)
    code, _out, err, pngs = render(deck(tmp_path / "d.pptx", hidden=()))
    assert code == 0
    assert "2 pages but 5 slides to render" in err
    assert pngs == {"s01.png": slide_width(1), "s02.png": slide_width(2)}


def test_a_deck_without_hidden_slides(tmp_path, libreoffice, render):
    libreoffice()
    code, out, _err, pngs = render(deck(tmp_path / "d.pptx", hidden=()))
    assert code == 0
    assert len(pngs) == 5
    assert "no hidden slides, so page N is slide N" in out


# ---------------------------------------------------------------- the font gate
def font_deck(path, runs):
    """One slide with one run per (family, bold)."""
    prs = Presentation()
    tf = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(6), Inches(1)).text_frame
    for family, bold in runs:
        r = tf.paragraphs[0].add_run()
        r.text, r.font.name, r.font.bold = "%s %s " % (family, bold), family, bold
    prs.save(str(path))
    return str(path)


def test_families_keeps_the_weight(tmp_path):
    prs = Presentation(font_deck(tmp_path / "d.pptx",
                                 [("Mono", True), ("Mono", False), ("Sans", True)]))
    assert sorted(R.families(prs)) == [("Mono", False), ("Mono", True), ("Sans", True)]


def test_a_missing_bold_warns_and_renders(tmp_path, fonts, libreoffice, render):
    fonts("Mono-hash.ttf", "Mono", "Regular")
    libreoffice()
    code, _out, err, pngs = render(font_deck(tmp_path / "d.pptx", [("Mono", True), ("Mono", False)]))
    assert code == 0
    assert pngs == {"s01.png": slide_width(1)}
    assert 'WARNING: "Mono" has no Bold file' in err
    assert "LibreOffice will synthesise the bold" in err


def test_a_missing_regular_still_stops_the_render(tmp_path, fonts, libreoffice, render):
    fonts("Mono-hash.ttf", "Mono", "Bold")
    libreoffice()
    code, _out, err, pngs = render(font_deck(tmp_path / "d.pptx", [("Mono", False)]))
    assert code == 2
    assert "the deck asks for Mono and no file for it was found" in err
    assert pngs == {}


def test_a_bold_only_family_with_no_file_at_all_stops_the_render(tmp_path, fonts, libreoffice,
                                                                render):
    libreoffice()
    code, _out, err, _pngs = render(font_deck(tmp_path / "d.pptx", [("Mono", True)]))
    assert code == 2
    assert "the deck asks for Mono and no file" in err


# ---------------------------------------------------------------- real LibreOffice
@pytest.mark.libreoffice
@pytest.mark.skipif(R.find_soffice() is None, reason="LibreOffice is not installed")
@pytest.mark.parametrize("flags, pages", [
    pytest.param((), ["SLIDE 1", "SLIDE 2 HIDDEN", "SLIDE 3", "SLIDE 4", "SLIDE 5 HIDDEN"],
                 id="rendered"),
    pytest.param(("--skip-hidden",), ["SLIDE 1", "SLIDE 3", "SLIDE 4"], id="skipped"),
])
def test_libreoffice_exports_what_the_flag_says(tmp_path, render, flags, pages):
    if "--skip-hidden" not in flags:
        version = subprocess.run([R.find_soffice(), "--version"],
                                 capture_output=True, text=True).stdout
        m = re.search(r"(\d+)\.(\d+)", version)
        if m and (int(m.group(1)), int(m.group(2))) < (7, 4):
            pytest.skip("LibreOffice %s.%s exports hidden slides from the command line "
                        "only from 7.4" % m.groups())
    code, _out, _err, pngs = render(deck(tmp_path / "d.pptx"), "--keep-pdf", *flags)
    assert code == 0
    doc = pymupdf.open(str(tmp_path / "render" / "d.pdf"))
    assert [page.get_text().strip() for page in doc] == pages
    assert len(pngs) == len(pages)
